"""Write the result sheet: one row per subject, plus a Run sheet recording how it was made.

The workbook has two sheets:

* **Subjects** - the rows, with the column contract from :mod:`core.export.columns`.
* **Run** - when the batch ran, which user asked, which sources answered, the codebook
  version and the per-status counts. It exists so a sheet that has been emailed on still
  answers "where did this come from?" without the original console output.

Excel-specific care taken here:

* identifier-shaped columns get the ``@`` (text) number format, so ``00177041`` does not
  come back as ``177041`` - the very corruption Tool 2 exists to find;
* dates are written as real dates with an ISO number format, so they sort correctly whatever
  the reader's locale;
* control characters that openpyxl refuses are stripped, and over-long texts are truncated
  at Excel's 32 767-character cell limit rather than raising.
"""

from __future__ import annotations

import logging
from collections.abc import Iterable, Mapping, Sequence
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Final

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.export.columns import (
    SUGGESTION_TEXT_COLUMNS,
    TEXT_COLUMNS,
    WRAPPED_COLUMNS,
    cell_value,
    header_label,
)

LOGGER = logging.getLogger(__name__)

SUBJECTS_SHEET: Final[str] = "Subjects"
RUN_SHEET: Final[str] = "Run"

#: Excel refuses a longer string in a single cell.
MAX_CELL_LENGTH: Final[int] = 32_767
_TRUNCATION_MARKER: Final[str] = " […]"

_DATE_FORMAT: Final[str] = "yyyy-mm-dd"
_DATETIME_FORMAT: Final[str] = "yyyy-mm-dd hh:mm:ss"
_TEXT_FORMAT: Final[str] = "@"

_HEADER_FILL: Final[PatternFill] = PatternFill("solid", fgColor="DDEBF7")
_HEADER_FONT: Final[Font] = Font(bold=True)

#: Column widths in characters: wrapped prose is wide, codes are narrow, the rest is default.
_WIDE: Final[int] = 60
_DEFAULT_WIDTH: Final[int] = 18
_MIN_WIDTH: Final[int] = 10
_MAX_MEASURED_WIDTH: Final[int] = 40


def _sanitize(text: str) -> str:
    """Strip characters openpyxl refuses and truncate at Excel's cell limit."""
    cleaned = ILLEGAL_CHARACTERS_RE.sub("", text)
    if len(cleaned) > MAX_CELL_LENGTH:
        LOGGER.warning("cell text of %d characters truncated to Excel's limit", len(cleaned))
        return cleaned[: MAX_CELL_LENGTH - len(_TRUNCATION_MARKER)] + _TRUNCATION_MARKER
    return cleaned


def _naive_utc(value: datetime) -> datetime:
    """Excel has no concept of a timezone, so store the UTC instant and drop the offset.

    Converting first matters: writing the wall-clock time of a non-UTC datetime unchanged
    would silently shift the timestamp. The affected headers say "(UTC)" so nobody reads the
    result as local time.
    """
    return (value.astimezone(UTC) if value.tzinfo is not None else value).replace(tzinfo=None)


def _write_cell(
    sheet: Worksheet, row_index: int, column_index: int, column: str, value: object
) -> None:
    """Write one value with the number format its column needs."""
    rendered = cell_value(column, value)
    if isinstance(rendered, str):
        rendered = _sanitize(rendered)
    elif isinstance(rendered, datetime):
        rendered = _naive_utc(rendered)
    cell = sheet.cell(row=row_index, column=column_index, value=rendered)

    if column in TEXT_COLUMNS or column in SUGGESTION_TEXT_COLUMNS:
        cell.number_format = _TEXT_FORMAT
    elif isinstance(rendered, datetime):
        cell.number_format = _DATETIME_FORMAT
    elif isinstance(rendered, date):
        cell.number_format = _DATE_FORMAT
    if column in WRAPPED_COLUMNS:
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _style_header(sheet: Worksheet, columns: Sequence[str]) -> None:
    """Bold, filled header with a freeze pane and an autofilter over the whole table."""
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=header_label(column))
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="top")
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 1:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(sheet.max_row, 1)}"


def _fit_columns(
    sheet: Worksheet, columns: Sequence[str], rows: Sequence[Mapping[str, object]]
) -> None:
    """Approximate column widths: wide for wrapped prose, measured for the rest."""
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        if column in WRAPPED_COLUMNS:
            sheet.column_dimensions[letter].width = _WIDE
            continue
        longest = len(header_label(column))
        for row in rows:
            rendered = cell_value(column, row.get(column))
            if rendered is not None:
                longest = max(longest, len(str(rendered)))
        sheet.column_dimensions[letter].width = min(
            max(longest + 2, _MIN_WIDTH), _MAX_MEASURED_WIDTH
        )


def _write_subjects(
    sheet: Worksheet, columns: Sequence[str], rows: Sequence[Mapping[str, object]]
) -> None:
    for row_index, row in enumerate(rows, start=2):
        for column_index, column in enumerate(columns, start=1):
            _write_cell(sheet, row_index, column_index, column, row.get(column))
    _style_header(sheet, columns)
    _fit_columns(sheet, columns, rows)


def _write_run(sheet: Worksheet, metadata: Mapping[str, object]) -> None:
    """Key/value sheet with the provenance of the whole run."""
    sheet.cell(row=1, column=1, value="key").font = _HEADER_FONT
    sheet.cell(row=1, column=2, value="value").font = _HEADER_FONT
    for index, (key, value) in enumerate(metadata.items(), start=2):
        sheet.cell(row=index, column=1, value=key)
        rendered = value
        if isinstance(rendered, (list, tuple)):
            rendered = ", ".join(str(item) for item in rendered) or None
        if isinstance(rendered, str):
            rendered = _sanitize(rendered)
        elif isinstance(rendered, datetime):
            rendered = _naive_utc(rendered)
        cell = sheet.cell(row=index, column=2, value=rendered)
        if isinstance(rendered, datetime):
            cell.number_format = _DATETIME_FORMAT
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = _WIDE
    sheet.freeze_panes = "A2"


def write_workbook(
    path: Path,
    columns: Sequence[str],
    rows: Iterable[Mapping[str, object]],
    *,
    run_metadata: Mapping[str, object] | None = None,
) -> Path:
    """Write the result workbook to ``path`` and return it.

    Args:
        path: Destination ``.xlsx``. Parent directories are created.
        columns: Column keys in sheet order (``IN_*`` first, then
            :data:`~core.export.columns.SUBJECT_COLUMNS`).
        rows: One mapping per subject, keyed by column. Missing keys become empty cells.
        run_metadata: Key/value pairs for the Run sheet; omitted when ``None``.
    """
    materialized = list(rows)
    workbook = Workbook()
    subjects = workbook.active
    subjects.title = SUBJECTS_SHEET
    _write_subjects(subjects, columns, materialized)
    if run_metadata:
        _write_run(workbook.create_sheet(RUN_SHEET), run_metadata)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    LOGGER.info("wrote %d row(s) to %s", len(materialized), path)
    return path
