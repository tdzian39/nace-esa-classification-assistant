"""Minimal, tolerant xlsx table reader built on openpyxl only.

Bank exports are messy: title rows above the header, headers spelled ``Kód`` / ``Kod`` /
``KÓD``, numbers stored as floats, trailing blank rows, non-breaking spaces. This module
turns such a sheet into plain ``dict[str, str]`` rows keyed by *logical* column names so the
loaders never touch openpyxl themselves and can later be swapped for another source.
"""

from __future__ import annotations

import unicodedata
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from core.codebooks.errors import CodebookFileError, CodebookSchemaError
from core.codebooks.normalize import normalize_cell


@dataclass(frozen=True)
class Table:
    """The usable content of one worksheet.

    Attributes:
        rows: One dict per non-blank data row, keyed by logical column name; every value has
            passed through :func:`~core.codebooks.normalize.normalize_cell`.
        row_numbers: The 1-based Excel row number of each entry in ``rows`` (same length), so
            problems can be reported by the number the user sees in Excel.
        sheet: Title of the worksheet that was read.
        header_row: 1-based row number of the detected header row.
        columns: Logical column name -> header text actually found in the file.
        path: The workbook that was read.
    """

    rows: tuple[dict[str, str], ...]
    row_numbers: tuple[int, ...]
    sheet: str
    header_row: int
    columns: dict[str, str]
    path: Path

    def __post_init__(self) -> None:
        if len(self.rows) != len(self.row_numbers):
            raise ValueError("rows and row_numbers must have the same length")

    def iter_numbered(self) -> Iterator[tuple[int, dict[str, str]]]:
        """Yield ``(excel_row_number, row)`` pairs in file order."""
        return zip(self.row_numbers, self.rows, strict=True)


def fold_header(text: object) -> str:
    """Canonical form used to compare header cells with aliases.

    Case-insensitive (``casefold``), whitespace-trimmed with inner runs collapsed, and
    accent-insensitive: NFKD decomposition with combining marks dropped, so ``"Kód"``,
    ``"Kod"`` and ``"KÓD"`` all fold to ``"kod"``.
    """
    decomposed = unicodedata.normalize("NFKD", normalize_cell(text))
    stripped = "".join(ch for ch in decomposed if not unicodedata.combining(ch))
    return " ".join(stripped.split()).casefold()


def _cell(row: Sequence[object], index: int) -> object:
    """Cell at ``index`` or ``None`` when the (read-only, unpadded) row is shorter."""
    return row[index] if index < len(row) else None


def _match_header(
    row: Sequence[object], columns: Mapping[str, Sequence[str]]
) -> tuple[dict[str, int], dict[str, str]] | None:
    """Try to interpret ``row`` as the header row.

    Returns ``(logical -> column index, logical -> header text)`` when every logical column
    is found on a distinct physical column, otherwise ``None``.
    """
    folded: dict[str, tuple[int, str]] = {}
    for index, value in enumerate(row):
        text = normalize_cell(value)
        if text:
            folded.setdefault(fold_header(text), (index, text))
    indexes: dict[str, int] = {}
    texts: dict[str, str] = {}
    for logical, aliases in columns.items():
        for alias in aliases:
            hit = folded.get(fold_header(alias))
            if hit is not None:
                indexes[logical] = hit[0]
                texts[logical] = hit[1]
                break
        else:
            return None
    if len(set(indexes.values())) != len(indexes):
        return None
    return indexes, texts


def read_table(
    path: Path,
    columns: Mapping[str, Sequence[str]],
    *,
    sheet: str | None = None,
    header_scan_rows: int = 20,
) -> Table:
    """Read the worksheet ``sheet`` (default: the first one) of ``path`` into a :class:`Table`.

    Args:
        path: The xlsx workbook.
        columns: Logical column name -> accepted header aliases, matched with
            :func:`fold_header` (case-, whitespace- and accent-insensitive).
        sheet: Worksheet title; ``None`` selects the first worksheet.
        header_scan_rows: The header is the first row within this many rows whose cells match
            ALL logical columns; rows above it (titles, export dates) are ignored.

    Rows where every logical column is empty are skipped, extra columns are ignored and the
    workbook is always closed.

    Raises:
        CodebookFileError: the file is missing, or openpyxl cannot open or read it (a
            corrupted sheet only fails while its rows are streamed in read-only mode).
        CodebookSchemaError: the sheet does not exist or no header row matches; the message
            lists the file, the sheet and the header texts that were found.
    """
    if not columns:
        raise ValueError("at least one logical column is required")
    path = Path(path)
    if not path.is_file():
        raise CodebookFileError(f"Codebook file not found: {path}")
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises BadZipFile, InvalidFileException, KeyError, ...
        raise CodebookFileError(f"Cannot open codebook file {path} as xlsx: {exc}") from exc
    try:
        worksheet = _select_worksheet(workbook, path, sheet)
        # Exports frequently carry a wrong <dimension> element; without this openpyxl would
        # truncate rows to the declared width in read-only mode.
        worksheet.reset_dimensions()
        header_row, indexes, texts = _find_header(worksheet, path, columns, header_scan_rows)
        rows: list[dict[str, str]] = []
        numbers: list[int] = []
        data = worksheet.iter_rows(min_row=header_row + 1, values_only=True)
        for number, raw in enumerate(data, start=header_row + 1):
            record = {
                logical: normalize_cell(_cell(raw, index)) for logical, index in indexes.items()
            }
            if any(record.values()):
                rows.append(record)
                numbers.append(number)
    except (CodebookFileError, CodebookSchemaError):
        raise
    except Exception as exc:  # truncated/garbled sheet XML surfaces lazily in read-only mode
        raise CodebookFileError(f"Cannot read codebook file {path} as xlsx: {exc}") from exc
    finally:
        workbook.close()
    return Table(
        rows=tuple(rows),
        row_numbers=tuple(numbers),
        sheet=worksheet.title,
        header_row=header_row,
        columns=texts,
        path=path,
    )


def _select_worksheet(
    workbook: openpyxl.Workbook, path: Path, sheet: str | None
) -> ReadOnlyWorksheet:
    """Return the named worksheet or the first one; raise ``CodebookSchemaError`` otherwise."""
    if sheet is not None:
        try:
            return workbook[sheet]
        except KeyError as exc:
            raise CodebookSchemaError(
                f"{path.name}: worksheet {sheet!r} not found; available: {workbook.sheetnames}"
            ) from exc
    worksheets = workbook.worksheets
    if not worksheets:
        raise CodebookSchemaError(f"{path.name}: workbook contains no worksheet")
    return worksheets[0]


def _find_header(
    worksheet: ReadOnlyWorksheet,
    path: Path,
    columns: Mapping[str, Sequence[str]],
    header_scan_rows: int,
) -> tuple[int, dict[str, int], dict[str, str]]:
    """Locate the header row; raise ``CodebookSchemaError`` listing what was found instead."""
    scan_rows = max(1, header_scan_rows)
    found_headers: list[str] = []
    rows = worksheet.iter_rows(min_row=1, max_row=scan_rows, values_only=True)
    for number, raw in enumerate(rows, start=1):
        matched = _match_header(raw, columns)
        if matched is not None:
            indexes, texts = matched
            return number, indexes, texts
        found_headers.extend(text for text in (normalize_cell(v) for v in raw) if text)
    title = worksheet.title
    expected = ", ".join(
        f"{logical} ({' | '.join(aliases)})" for logical, aliases in columns.items()
    )
    raise CodebookSchemaError(
        f"{path.name} sheet {title!r}: no header row with columns [{expected}] within the "
        f"first {scan_rows} rows; headers found: {found_headers or 'none'}"
    )
