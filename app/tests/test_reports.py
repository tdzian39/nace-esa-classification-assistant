"""Error reports: the report, the two stores, the review CLI, and the page's button.

Nothing here touches the network: the Blob store gets an ``httpx.MockTransport``.
"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api import main as api
from config.settings import Settings
from core.classify.provider import StubLlmProvider
from core.reports import (
    BlobReportStore,
    DirectoryReportStore,
    ErrorReport,
    NullReportStore,
    ReportStoreError,
    _main,
    build_report,
    build_report_store,
    reports_are_volatile,
    review_row,
)
from tests.api.test_api_suggest import answer, make_service
from tests.classify.conftest import build_codebooks

NOW = datetime(2026, 9, 24, 12, 0, tzinfo=UTC)


def suggestion():
    codebooks = build_codebooks()
    provider = StubLlmProvider({"NACE": answer("64"), "ESA": answer("2002703")})
    from core.suggest import SuggestionRequest

    return make_service(codebooks, provider=provider).suggest(
        SuggestionRequest(name="Nordkap Funding B.V.")
    )


def report(note: str = "  Je to  kaptivní  společnost. ", **kwargs) -> ErrorReport:
    return build_report(suggestion(), note=note, user="jana novak", now=NOW, **kwargs)


class TestReport:
    def test_it_carries_the_request_and_the_result_row(self) -> None:
        item = report()
        assert item.request == {"isin": None, "name": "Nordkap Funding B.V.", "description": None}
        assert item.result["NACE_code"] == "64"
        assert item.result["ESA_code"] == "2002703"
        assert item.result["IN_name"] == "Nordkap Funding B.V."
        assert item.answered is True
        assert item.user == "jana novak"
        assert item.created_at == NOW

    def test_the_note_is_trimmed_and_capped(self) -> None:
        assert report().note == "Je to kaptivní společnost."
        assert len(report("x" * 900, max_note_chars=10).note) == 10

    def test_the_app_facts_travel_with_it(self) -> None:
        item = report(commit="abc123")
        assert item.app["commit"] == "abc123"
        assert item.app["codebook_version"]
        assert item.app["model"]

    def test_json_round_trip(self) -> None:
        item = report()
        back = ErrorReport.from_dict(json.loads(item.to_json()))
        assert back == item

    def test_the_pathname_is_one_folder_per_day(self) -> None:
        item = report()
        assert item.pathname("reports/") == f"reports/2026-09-24/{item.id}.json"

    def test_the_identifier_for_the_audit_line(self) -> None:
        assert report().identifier == "Nordkap Funding B.V."


class TestDirectoryStore:
    def test_writes_one_json_file_per_report(self, tmp_path: Path) -> None:
        store = DirectoryReportStore(tmp_path / "reports")
        item = report()
        location = store.save(item)
        path = Path(location)
        assert path == tmp_path / "reports" / "2026-09-24" / f"{item.id}.json"
        assert json.loads(path.read_text(encoding="utf-8"))["kind"] == "error_report"
        assert not list(path.parent.glob(".report-*"))  # no temp file left behind

    def test_never_overwrites(self, tmp_path: Path) -> None:
        store = DirectoryReportStore(tmp_path)
        item = report()
        store.save(item)
        with pytest.raises(ReportStoreError):
            store.save(item)

    def test_an_unwritable_root_is_an_error_not_a_crash(self, tmp_path: Path) -> None:
        blocker = tmp_path / "file"
        blocker.write_text("not a directory")
        with pytest.raises(ReportStoreError):
            DirectoryReportStore(blocker / "reports").save(report())

    def test_load_reads_them_back_and_skips_strangers(self, tmp_path: Path) -> None:
        store = DirectoryReportStore(tmp_path)
        first = report("first")
        store.save(first)
        (tmp_path / "2026-09-24" / "junk.json").write_text("{not json", encoding="utf-8")
        (tmp_path / "2026-09-24" / "other.json").write_text('{"kind": "x"}', encoding="utf-8")
        loaded = list(store.load())
        assert loaded == [first]

    def test_a_missing_root_holds_nothing(self, tmp_path: Path) -> None:
        assert list(DirectoryReportStore(tmp_path / "nope").load()) == []


def blob_client(handler) -> httpx.Client:
    return httpx.Client(transport=httpx.MockTransport(handler))


class TestBlobStore:
    def store(self, handler, **kwargs) -> BlobReportStore:
        return BlobReportStore(
            token="vercel_blob_rw_abc123_secret",
            store_id="abc123",
            client=blob_client(handler),
            **kwargs,
        )

    def test_puts_the_json_the_way_the_sdk_does(self) -> None:
        seen: list[httpx.Request] = []

        def handler(request: httpx.Request) -> httpx.Response:
            seen.append(request)
            return httpx.Response(
                200, json={"url": "https://x", "pathname": request.url.params["pathname"]}
            )

        item = report()
        location = self.store(handler).save(item)
        (request,) = seen
        assert request.method == "PUT"
        assert str(request.url).startswith("https://vercel.com/api/blob/?pathname=")
        assert request.url.params["pathname"] == f"reports/2026-09-24/{item.id}.json"
        assert request.headers["authorization"] == "Bearer vercel_blob_rw_abc123_secret"
        assert request.headers["x-api-version"] == "12"
        assert request.headers["x-vercel-blob-access"] == "private"
        assert request.headers["x-vercel-blob-store-id"] == "abc123"
        assert request.headers["x-content-type"] == "application/json"
        assert request.headers["x-add-random-suffix"] == "0"
        assert request.headers["x-allow-overwrite"] == "0"
        assert json.loads(request.content)["id"] == item.id
        assert location == f"reports/2026-09-24/{item.id}.json"

    def test_a_5xx_is_retried_once_then_an_error(self) -> None:
        calls = 0

        def handler(request: httpx.Request) -> httpx.Response:
            nonlocal calls
            calls += 1
            return httpx.Response(503, text="boom")

        with pytest.raises(ReportStoreError, match="HTTP 503"):
            self.store(handler).save(report())
        assert calls == 2

    def test_a_4xx_is_not_retried(self) -> None:
        calls = 0

        def handler(request: httpx.Request) -> httpx.Response:
            nonlocal calls
            calls += 1
            return httpx.Response(403, text="nope")

        with pytest.raises(ReportStoreError, match="HTTP 403"):
            self.store(handler).save(report())
        assert calls == 1

    def test_unreachable_is_an_error(self) -> None:
        def handler(request: httpx.Request) -> httpx.Response:
            raise httpx.ConnectError("down")

        with pytest.raises(ReportStoreError, match="unreachable"):
            self.store(handler).save(report())

    def test_the_prefix_and_api_url_are_settings(self) -> None:
        seen: list[httpx.Request] = []

        def handler(request: httpx.Request) -> httpx.Response:
            seen.append(request)
            return httpx.Response(200, json={})

        self.store(handler, prefix="mo/", api_url="https://blob.example/api/").save(report())
        assert str(seen[0].url).startswith("https://blob.example/api/?pathname=mo%2F2026-09-24")


class TestBuildStore:
    def test_dir_is_the_default_and_relative_to_the_app(self) -> None:
        store = build_report_store(Settings())
        assert isinstance(store, DirectoryReportStore)
        assert store.root.is_absolute()
        assert store.root.name == "reports"

    def test_off_is_a_null_store(self) -> None:
        store = build_report_store(Settings(reports_source="off"))
        assert isinstance(store, NullReportStore)
        with pytest.raises(ReportStoreError):
            store.save(report())

    def test_blob_without_a_token_is_a_readable_error(self) -> None:
        with pytest.raises(ReportStoreError, match="BLOB_READ_WRITE_TOKEN"):
            build_report_store(Settings(reports_source="blob"))

    def test_blob_takes_the_store_id_from_the_token(self) -> None:
        store = build_report_store(
            Settings(reports_source="blob", blob_read_write_token="vercel_blob_rw_st0re_secret")
        )
        assert isinstance(store, BlobReportStore)

    def test_a_directory_on_vercel_is_volatile(self) -> None:
        assert reports_are_volatile(Settings(), {"VERCEL": "1"})
        assert not reports_are_volatile(Settings(), {})
        assert not reports_are_volatile(Settings(reports_source="blob"), {"VERCEL": "1"})


class TestReview:
    def test_the_row_is_flat(self) -> None:
        row = review_row(report(commit="abc"))
        assert row["NACE_code"] == "64"
        assert row["ESA_cts_id"]
        assert row["commit"] == "abc"
        assert row["created_at (UTC)"] == NOW.replace(tzinfo=None)

    def test_the_cli_lists_and_exports(self, tmp_path: Path, capsys: pytest.CaptureFixture) -> None:
        store = DirectoryReportStore(tmp_path / "r")
        store.save(report("první"))
        store.save(report("druhá"))
        assert _main(["--dir", str(tmp_path / "r"), "--list"]) == 0
        out = capsys.readouterr().out
        assert "2 report(s)" in out and "první" in out
        target = tmp_path / "reports.xlsx"
        assert _main(["--dir", str(tmp_path / "r"), "--xlsx", str(target)]) == 0
        sheet = load_workbook(target)
        assert any("Nordkap" in str(c.value) for ws in sheet for row in ws.iter_rows() for c in row)

    def test_the_cli_without_a_verb_prints_help(self) -> None:
        assert _main([]) == 2


@pytest.fixture
def client(tmp_path: Path):
    codebooks = build_codebooks()
    provider = StubLlmProvider({"NACE": answer("64"), "ESA": answer("2002703")})
    api._state["settings"] = Settings(llm_api_key=None, llm_cache_path=None)
    api._state["service"] = make_service(codebooks, provider=provider)
    api._state["reports"] = DirectoryReportStore(tmp_path / "reports")
    try:
        yield TestClient(api.app), tmp_path / "reports"
    finally:
        api._state.clear()


class TestPage:
    def test_the_result_offers_the_button(self, client) -> None:
        test_client, _ = client
        page = test_client.post("/suggest", data={"name": "Nordkap Funding B.V."}).text
        assert "Nahlásit k prověření" in page
        assert 'action="/report"' in page
        assert 'name="note"' in page

    def test_the_form_page_does_not(self, client) -> None:
        test_client, _ = client
        assert "Nahlásit k prověření" not in test_client.get("/").text

    def test_a_report_is_stored_with_the_request_and_the_result(self, client) -> None:
        test_client, root = client
        response = test_client.post(
            "/report",
            data={"name": "Nordkap Funding B.V.", "note": "Spíš kaptivní společnost."},
            headers={"X-Remote-User": "jana novak"},
        )
        assert response.status_code == 200
        assert "Hlášení uloženo" in response.text
        assert "NACE" in response.text  # the result is still on the page
        (path,) = list(root.rglob("*.json"))
        stored = json.loads(path.read_text(encoding="utf-8"))
        assert stored["request"]["name"] == "Nordkap Funding B.V."
        assert stored["result"]["NACE_code"] == "64"
        assert stored["note"] == "Spíš kaptivní společnost."
        assert stored["user"] == "jana novak"
        assert stored["app"]["codebook_version"]

    def test_the_note_may_be_empty(self, client) -> None:
        test_client, root = client
        response = test_client.post("/report", data={"name": "Nordkap", "note": ""})
        assert "Hlášení uloženo" in response.text
        (path,) = list(root.rglob("*.json"))
        assert json.loads(path.read_text(encoding="utf-8"))["note"] == ""

    def test_a_stored_report_shows_no_second_button(self, client) -> None:
        test_client, _ = client
        page = test_client.post("/report", data={"name": "Nordkap"}).text
        assert "Nahlásit k prověření" not in page

    def test_a_store_failure_is_said_on_the_page(self, client) -> None:
        test_client, _ = client
        api._state["reports"] = NullReportStore()
        response = test_client.post("/report", data={"name": "Nordkap", "note": "x"})
        assert response.status_code == 200
        assert "Hlášení se nepodařilo uložit" in response.text

    def test_an_empty_report_is_refused(self, client) -> None:
        test_client, root = client
        response = test_client.post("/report", data={"note": "x"})
        assert "Hlášení bez dotazu nelze uložit" in response.text
        assert not list(root.rglob("*.json"))

    def test_off_hides_the_button(self, client) -> None:
        test_client, _ = client
        api._state["settings"] = Settings(
            llm_api_key=None, llm_cache_path=None, reports_source="off"
        )
        page = test_client.post("/suggest", data={"name": "Nordkap"}).text
        assert "Nahlásit k prověření" not in page

    def test_the_audit_line_names_the_user_but_not_the_note(
        self, client, caplog: pytest.LogCaptureFixture
    ) -> None:
        test_client, _ = client
        with caplog.at_level("INFO", logger="core.audit"):
            test_client.post(
                "/report",
                data={"name": "Nordkap", "note": "TAJNÁ POZNÁMKA"},
                headers={"X-Remote-User": "jana novak"},
            )
        lines = [r.getMessage() for r in caplog.records if r.name == "core.audit"]
        assert any("report identifier='Nordkap' user=jana novak outcome=stored" in m for m in lines)
        assert not any("TAJNÁ" in m for m in lines)
