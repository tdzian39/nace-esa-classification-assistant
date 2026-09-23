"""The API and the page, driven through FastAPI's test client.

The app is wired with the real codebooks-to-suggestion pipeline and a stubbed model, so what
is exercised is the actual request handling - templates, forms, JSON, the xlsx download -
without a key or a network.
"""

from __future__ import annotations

import io
import json
from collections.abc import Iterator

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api import main as api
from config.settings import Settings
from core.classify.llm import LlmClassifier
from core.classify.provider import NullLlmProvider, StubLlmProvider
from core.codebooks.models import CodebookSet
from core.sources.gleif import GleifSource
from core.sources.identity import IssuerIdentifier
from core.sources.openfigi import OpenFigiSource
from core.sources.web import SearchHit, StaticSearchProvider, WebEvidenceGatherer
from core.suggest import SuggestionService
from tests.classify.conftest import build_codebooks
from tests.sources.conftest import (
    FIGI_DEUTSCHE_BANK,
    GLEIF_DEUTSCHE_BANK,
    figi_client,
    gleif_client,
)

PAGE = (
    "<html><head><meta name='description' content='Nordkap Funding B.V. is the financing "
    "vehicle of the Nordkap group.'></head><body><p>The company issues bonds and on-lends "
    "the proceeds exclusively to its parent bank. It holds no banking licence.</p>"
    "</body></html>"
)


def answer(code: str, confidence: str = "high", why: str = "Financuje vlastní skupinu.") -> str:
    return json.dumps(
        {
            "sufficient_evidence": True,
            "picks": [{"code": code, "confidence": confidence, "justification": why}],
        }
    )


def make_service(
    codebooks: CodebookSet,
    *,
    provider=None,
    hits: tuple[SearchHit, ...] = (
        SearchHit(url="https://example.com/a", title="Nordkap Funding B.V."),
    ),
    identifier: IssuerIdentifier | None = None,
) -> SuggestionService:
    import httpx

    settings = Settings(web_min_interval_seconds=0.0, llm_cache_path=None, llm_api_key=None)
    client = httpx.Client(
        transport=httpx.MockTransport(
            lambda request: httpx.Response(200, text=PAGE, headers={"content-type": "text/html"})
        )
    )
    gatherer = WebEvidenceGatherer(
        settings, provider=StaticSearchProvider(hits), client=client, sleep=lambda _: None
    )
    return SuggestionService(
        codebooks,
        gatherer=gatherer,
        classifier=LlmClassifier(provider or NullLlmProvider()),
        identifier=identifier,
    )


@pytest.fixture
def client(monkeypatch: pytest.MonkeyPatch) -> Iterator[TestClient]:
    """A client whose app is already wired, bypassing the codebook-loading lifespan."""
    codebooks = build_codebooks()
    provider = StubLlmProvider({"NACE": answer("64"), "ESA": answer("2002703")})
    api._state["settings"] = Settings(llm_api_key=None, llm_cache_path=None)
    api._state["service"] = make_service(codebooks, provider=provider)
    try:
        yield TestClient(api.app)
    finally:
        api._state.clear()


class TestPage:
    def test_the_form_is_served(self, client: TestClient) -> None:
        response = client.get("/")
        assert response.status_code == 200
        assert 'name="isin"' in response.text
        assert 'name="description"' in response.text

    def test_a_lookup_renders_both_panels(self, client: TestClient) -> None:
        response = client.post("/suggest", data={"name": "Nordkap Funding B.V."})
        assert response.status_code == 200
        assert "NACE" in response.text and "ESA 2010" in response.text

    def test_the_cts_id_is_shown_next_to_the_code(self, client: TestClient) -> None:
        """It is the number MO types into CTS, so it has to be on screen, not implied."""
        response = client.post("/suggest", data={"name": "Nordkap Funding B.V."})
        assert "CTS ID" in response.text
        assert "64" in response.text

    def test_the_justification_is_shown(self, client: TestClient) -> None:
        response = client.post("/suggest", data={"name": "Nordkap"})
        assert "Financuje vlastní skupinu." in response.text

    def test_the_form_keeps_what_was_typed(self, client: TestClient) -> None:
        response = client.post("/suggest", data={"name": "Nordkap Funding B.V."})
        assert 'value="Nordkap Funding B.V."' in response.text

    def test_an_empty_form_is_rejected_without_a_lookup(self, client: TestClient) -> None:
        response = client.post("/suggest", data={"isin": "", "name": "", "description": ""})
        assert response.status_code == 200
        assert "Vyplňte alespoň jednu položku" in response.text

    def test_evidence_links_are_offered(self, client: TestClient) -> None:
        response = client.post("/suggest", data={"name": "Nordkap"})
        assert "Podklady" in response.text
        assert "https://example.com/a" in response.text

    def test_a_warning_says_when_the_model_is_not_configured(self, client: TestClient) -> None:
        """The page must say the tool is only shortlisting, not choosing."""
        assert "LLM_API_KEY" in client.get("/").text


#: A typed description no keyword and no register rule speaks to: winegrowing.
UNRULED = "Pěstování vinné révy a výroba vína v údolí Rýna."


class TestAbstention:
    def test_an_abstention_is_shown_not_hidden(self, monkeypatch: pytest.MonkeyPatch) -> None:
        """A blank panel would look like a bug; the reason has to be on screen."""
        codebooks = build_codebooks()
        api._state["settings"] = Settings(llm_api_key=None, llm_cache_path=None)
        api._state["service"] = make_service(codebooks, provider=NullLlmProvider())
        try:
            response = TestClient(api.app).post("/suggest", data={"description": UNRULED})
            assert "Nástroj kód nevybral" in response.text
            assert "navrhovaný kód" not in response.text
            assert "no model configured" in response.text
            # Deterministic mode is a result, not a failure: the narrowed codebook has to be
            # on screen with its CTS IDs, or the tool has thrown its own work away.
            assert "zúžený číselník" in response.text
            assert "CTS ID" in response.text
        finally:
            api._state.clear()

    def test_without_a_model_a_rule_s_code_is_proposed_and_says_so(self) -> None:
        """The brief asks for a suggested code; a rule's pick is one, labelled as the rules'."""
        codebooks = build_codebooks()
        api._state["settings"] = Settings(llm_api_key=None, llm_cache_path=None)
        api._state["service"] = make_service(codebooks, provider=NullLlmProvider())
        try:
            response = TestClient(api.app).post("/suggest", data={"name": "Nordkap Funding B.V."})
            assert "navrhovaný kód" in response.text
            assert "podle pravidel" in response.text
            assert "Podle pravidel, bez modelu" in response.text
            assert "jistota" not in response.text  # only a model has a confidence
            assert "no model configured" in response.text  # why the rules had to decide
            # The rest of the narrowed codebook stays on screen under the proposal.
            assert "další kandidáti ze zúženého číselníku" in response.text
        finally:
            api._state.clear()

    def test_the_json_names_the_basis_of_each_proposal(self) -> None:
        codebooks = build_codebooks()
        api._state["settings"] = Settings(llm_api_key=None, llm_cache_path=None)
        api._state["service"] = make_service(codebooks, provider=NullLlmProvider())
        try:
            client = TestClient(api.app)
            ruled = client.post("/api/suggest", json={"name": "Nordkap Funding B.V."}).json()
            assert ruled["answered"] is False  # the model did not answer ...
            assert ruled["nace"]["proposal"]["basis"] == "rules"  # ... a rule did
            assert ruled["nace"]["proposal"]["code"] == "64"
            assert ruled["nace"]["proposal"]["confidence"] is None
            assert ruled["row"]["NACE_code"] == "64"
            assert ruled["row"]["NACE_confidence"] is None
            assert ruled["row"]["NACE_justification"].startswith("Podle pravidel, bez modelu")

            unruled = client.post("/api/suggest", json={"description": UNRULED}).json()
            assert unruled["nace"]["proposal"] is None
            assert unruled["row"]["NACE_code"] is None
        finally:
            api._state.clear()

    def test_no_description_still_returns_a_page(self, monkeypatch: pytest.MonkeyPatch) -> None:
        codebooks = build_codebooks()
        api._state["settings"] = Settings(llm_api_key=None, llm_cache_path=None)
        api._state["service"] = make_service(codebooks, provider=NullLlmProvider(), hits=())
        try:
            response = TestClient(api.app).post("/suggest", data={"name": "Neznámá firma"})
            assert response.status_code == 200
            assert "nepodařilo získat" in response.text
        finally:
            api._state.clear()


class TestJson:
    def test_a_lookup_returns_the_row_and_both_classifications(self, client: TestClient) -> None:
        response = client.post("/api/suggest", json={"name": "Nordkap Funding B.V."})
        assert response.status_code == 200
        body = response.json()

        assert body["answered"] is True
        assert body["nace"]["suggestions"][0]["code"] == "64"
        assert body["esa"]["suggestions"][0]["code"] == "2002703"
        assert body["nace"]["proposal"]["basis"] == "model"
        assert body["esa"]["proposal"]["code"] == "2002703"
        assert body["row"]["NACE_cts_id"]

    def test_an_empty_request_is_422(self, client: TestClient) -> None:
        assert client.post("/api/suggest", json={}).status_code == 422

    def test_the_row_carries_provenance(self, client: TestClient) -> None:
        row = client.post("/api/suggest", json={"name": "Nordkap"}).json()["row"]
        assert row["source"] == "WEB"
        assert row["retrieved_at"]
        assert row["codebook_version"]

    def test_notes_explain_a_malformed_isin(self, client: TestClient) -> None:
        """A bad ISIN is a note, not a rejection: the name may still be enough."""
        body = client.post("/api/suggest", json={"isin": "NOTANISIN", "name": "Nordkap"}).json()
        assert any("ISIN" in note for note in body["notes"])
        assert body["answered"] is True


class TestDownload:
    def test_the_xlsx_opens_and_holds_the_row(self, client: TestClient) -> None:
        response = client.get("/suggest.xlsx", params={"name": "Nordkap Funding B.V."})
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["content-type"]

        book = load_workbook(io.BytesIO(response.content))
        sheet = book["Subjects"]
        headers = [cell.value for cell in sheet[1]]
        assert "NACE_cts_id (do CTS)" in headers
        assert sheet.max_row == 2

    def test_codes_are_written_as_text(self, client: TestClient) -> None:
        """A NACE code losing its leading zero on the way out would be the tool's own bug."""
        response = client.get("/suggest.xlsx", params={"name": "Nordkap"})
        sheet = load_workbook(io.BytesIO(response.content))["Subjects"]
        headers = [cell.value for cell in sheet[1]]
        column = headers.index("NACE_code") + 1
        assert sheet.cell(row=2, column=column).number_format == "@"

    def test_the_run_sheet_records_provenance(self, client: TestClient) -> None:
        response = client.get("/suggest.xlsx", params={"name": "Nordkap"})
        book = load_workbook(io.BytesIO(response.content))
        assert "Run" in book.sheetnames
        pairs = {
            book["Run"].cell(row=r, column=1).value: book["Run"].cell(row=r, column=2).value
            for r in range(2, book["Run"].max_row + 1)
        }
        assert pairs["nástroj"] == "ESA a NACE našeptávač"
        assert pairs["verze číselníku"]

    def test_the_filename_follows_the_issuer(self, client: TestClient) -> None:
        response = client.get("/suggest.xlsx", params={"name": "Nordkap Funding B.V."})
        assert "Nordkap" in response.headers["content-disposition"]

    def test_an_empty_request_yields_an_empty_sheet_not_an_error(self, client: TestClient) -> None:
        response = client.get("/suggest.xlsx")
        assert response.status_code == 200
        assert load_workbook(io.BytesIO(response.content))["Subjects"].max_row == 1


class TestHealth:
    def test_health_reports_what_an_operator_needs(self, client: TestClient) -> None:
        body = client.get("/health").json()
        assert body["status"] == "ok"
        assert body["codebook_version"]
        assert body["llm_configured"] is False
        assert body["search_configured"] is False


def test_every_lookup_is_audited(client: TestClient, caplog: pytest.LogCaptureFixture) -> None:
    with caplog.at_level("INFO", logger="core.audit"):
        client.post("/suggest", data={"name": "Nordkap Funding B.V."})
    assert any("lookup identifier=" in record.message for record in caplog.records)


class TestAuditedUser:
    """Who a lookup is recorded against.

    The brief requires the *requesting* user to be logged. In a server the OS account is the
    service account, so without a trusted header every lookup would be attributed to the same
    name - an audit trail that cannot answer who asked.
    """

    def test_the_signed_in_user_is_recorded(
        self, client: TestClient, caplog: pytest.LogCaptureFixture
    ) -> None:
        with caplog.at_level("INFO", logger="core.audit"):
            client.post(
                "/suggest",
                data={"name": "Nordkap"},
                headers={"X-Remote-User": "mo.analyst"},
            )
        assert any("user=mo.analyst" in record.message for record in caplog.records)

    def test_the_json_endpoint_records_it_too(
        self, client: TestClient, caplog: pytest.LogCaptureFixture
    ) -> None:
        with caplog.at_level("INFO", logger="core.audit"):
            client.post(
                "/api/suggest", json={"name": "Nordkap"}, headers={"X-Remote-User": "reporting.jk"}
            )
        assert any("user=reporting.jk" in record.message for record in caplog.records)

    def test_the_download_records_it_too(
        self, client: TestClient, caplog: pytest.LogCaptureFixture
    ) -> None:
        with caplog.at_level("INFO", logger="core.audit"):
            client.get(
                "/suggest.xlsx", params={"name": "Nordkap"}, headers={"X-Remote-User": "mo.two"}
            )
        assert any("user=mo.two" in record.message for record in caplog.records)

    def test_the_header_name_is_configurable(
        self, client: TestClient, caplog: pytest.LogCaptureFixture
    ) -> None:
        api._state["settings"] = Settings(
            llm_api_key=None, llm_cache_path=None, web_user_header="X-Forwarded-User"
        )
        with caplog.at_level("INFO", logger="core.audit"):
            client.post(
                "/suggest", data={"name": "Nordkap"}, headers={"X-Forwarded-User": "sso.user"}
            )
        assert any("user=sso.user" in record.message for record in caplog.records)

    def test_without_the_header_it_falls_back_rather_than_failing(
        self, client: TestClient, caplog: pytest.LogCaptureFixture
    ) -> None:
        """A local run has no proxy; the OS account really is the person there."""
        with caplog.at_level("INFO", logger="core.audit"):
            response = client.post("/suggest", data={"name": "Nordkap"})
        assert response.status_code == 200
        assert any("user=" in record.message for record in caplog.records)

    def test_the_download_names_the_user_in_the_run_sheet(self, client: TestClient) -> None:
        response = client.get(
            "/suggest.xlsx", params={"name": "Nordkap"}, headers={"X-Remote-User": "mo.analyst"}
        )
        book = load_workbook(io.BytesIO(response.content))
        pairs = {
            book["Run"].cell(row=r, column=1).value: book["Run"].cell(row=r, column=2).value
            for r in range(2, book["Run"].max_row + 1)
        }
        assert pairs["uživatel"] == "mo.analyst"


class TestIsinIdentity:
    """An ISIN is resolved to its issuer through GLEIF and OpenFIGI before anything is searched.

    The registers are driven through ``httpx.MockTransport`` clients answering with trimmed
    copies of the live payloads for Deutsche Bank AG (DE0005140008), so the whole chain -
    identifier, evidence, shortlist, classifier, page, JSON, audit - runs without a network.
    """

    ISIN = "DE0005140008"
    LEI = "7LTWFZYICNSX8D621K86"

    @pytest.fixture
    def calls(self) -> list[str]:
        return []

    @pytest.fixture
    def isin_client(self, calls: list[str]) -> Iterator[TestClient]:
        settings = Settings(
            llm_api_key=None,
            llm_cache_path=None,
            gleif_min_interval_seconds=0.0,
            gleif_max_attempts=1,
            openfigi_min_interval_seconds=0.0,
            openfigi_max_attempts=1,
        )
        identifier = IssuerIdentifier(
            settings,
            gleif=GleifSource(
                settings,
                client=gleif_client(
                    by_isin={self.ISIN: GLEIF_DEUTSCHE_BANK},
                    exceptions={self.LEI: "NO_KNOWN_PERSON"},
                    calls=calls,
                ),
                sleep=lambda _: None,
            ),
            openfigi=OpenFigiSource(
                settings, client=figi_client(FIGI_DEUTSCHE_BANK, calls=calls), sleep=lambda _: None
            ),
        )
        provider = StubLlmProvider({"NACE": answer("64"), "ESA": answer("2002213")})
        api._state["settings"] = settings
        api._state["service"] = make_service(
            build_codebooks(), provider=provider, hits=(), identifier=identifier
        )
        try:
            yield TestClient(api.app)
        finally:
            api._state.clear()

    def test_the_issuer_is_named_from_the_register(self, isin_client: TestClient) -> None:
        """An ISIN alone used to yield 'Emitent neurčen'; now the LEI record names it."""
        response = isin_client.post("/suggest", data={"isin": self.ISIN})
        assert response.status_code == 200
        assert "DEUTSCHE BANK AKTIENGESELLSCHAFT" in response.text
        assert f"LEI {self.LEI}" in response.text

    def test_the_register_facts_are_on_the_page(self, isin_client: TestClient) -> None:
        text = isin_client.post("/suggest", data={"isin": self.ISIN}).text
        assert "Kategorie subjektu podle GLEIF" in text
        assert "OpenFIGI (ISIN DE0005140008)" in text
        assert "Mateřská společnost není v GLEIF uvedena" in text

    def test_the_register_pages_are_cited(self, isin_client: TestClient) -> None:
        text = isin_client.post("/suggest", data={"isin": self.ISIN}).text
        assert "search.gleif.org" in text
        assert "openfigi.com" in text

    def test_the_facts_reach_the_classifier(self, isin_client: TestClient) -> None:
        """The legal name carries 'BANK', so the keyword hints put 64 and the bank family on."""
        body = isin_client.post("/api/suggest", json={"isin": self.ISIN}).json()
        assert body["answered"] is True
        assert body["nace"]["suggestions"][0]["code"] == "64"
        assert body["esa"]["suggestions"][0]["code"] == "2002213"

    def test_json_carries_the_identity(self, isin_client: TestClient) -> None:
        body = isin_client.post("/api/suggest", json={"isin": self.ISIN}).json()
        assert body["issuer_name"] == "DEUTSCHE BANK AKTIENGESELLSCHAFT"
        assert body["identity"]["lei"] == self.LEI
        assert body["identity"]["country"] == "DE"
        assert body["identity"]["category"] == "GENERAL"
        assert body["identity"]["instrument"]["market_sector"] == "Equity"
        assert body["identity"]["sources"] == ["GLEIF", "OPENFIGI"]

    def test_the_row_names_its_sources_and_the_lei(self, isin_client: TestClient) -> None:
        row = isin_client.post("/api/suggest", json={"isin": self.ISIN}).json()["row"]
        assert row["source"] == "GLEIF+OPENFIGI+WEB"
        assert row["issuer_lei"] == self.LEI
        assert row["issuer_country"] == "DE"
        assert any("gleif.org" in url for url in row["evidence_urls"])

    def test_the_audit_names_the_registers(
        self, isin_client: TestClient, caplog: pytest.LogCaptureFixture
    ) -> None:
        with caplog.at_level("INFO", logger="core.audit"):
            isin_client.post("/suggest", data={"isin": self.ISIN})
        assert any("sources=GLEIF+OPENFIGI+WEB" in record.message for record in caplog.records)

    def test_a_typed_name_still_wins(self, isin_client: TestClient) -> None:
        """What MO typed is authoritative; the register's spelling is shown as a fact."""
        body = isin_client.post(
            "/api/suggest", json={"isin": self.ISIN, "name": "Deutsche Bank AG"}
        ).json()
        assert body["issuer_name"] == "Deutsche Bank AG"
        assert body["identity"]["legal_name"] == "DEUTSCHE BANK AKTIENGESELLSCHAFT"

    def test_a_name_lookup_does_not_touch_the_registers(
        self, isin_client: TestClient, calls: list[str]
    ) -> None:
        isin_client.post("/suggest", data={"name": "Nordkap Funding B.V."})
        assert calls == []

    def test_the_download_carries_the_lei(self, isin_client: TestClient) -> None:
        response = isin_client.get("/suggest.xlsx", params={"isin": self.ISIN})
        sheet = load_workbook(io.BytesIO(response.content))["Subjects"]
        headers = [cell.value for cell in sheet[1]]
        column = headers.index("issuer_lei (GLEIF)") + 1
        assert sheet.cell(row=2, column=column).value == self.LEI

    def test_health_reports_the_registers(self, isin_client: TestClient) -> None:
        body = isin_client.get("/health").json()
        assert body["gleif_enabled"] is True
        assert body["openfigi_enabled"] is True
