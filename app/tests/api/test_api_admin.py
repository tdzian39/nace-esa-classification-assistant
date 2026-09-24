"""The developer page: its own password, the priced ledger, the complaints, the downloads."""

from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import UTC, date, datetime, timedelta
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api import main as api
from config.settings import Settings
from core.admin import LedgerFilter, ReportFilter, ledger_view, parse_day, reports_view
from core.auth import hash_password
from core.classify.budget import UsageRecord
from core.classify.provider import StubLlmProvider
from core.db import Database, DatabaseAuditSink, DatabaseLedger, DatabaseReportStore
from core.reports import build_report
from tests.api.test_api_suggest import answer, make_service
from tests.classify.conftest import build_codebooks

ADMIN_PASSWORD = "dev-secret"
ADMIN_HASH = hash_password(ADMIN_PASSWORD, iterations=1000)
MO_PASSWORD = "mo-secret"
MO_HASH = hash_password(MO_PASSWORD, iterations=1000)
NOW = datetime(2026, 9, 24, 12, 0, tzinfo=UTC)


def record(**overrides) -> UsageRecord:
    base = {
        "at": NOW,
        "model": "gpt-5.6-luna",
        "kind": "NACE",
        "prompt_tokens": 1000,
        "completion_tokens": 100,
        "cached_prompt_tokens": 500,
        "user": "jana",
    }
    base.update(overrides)
    return UsageRecord(**base)


class TestLedgerView:
    def test_totals_and_groups(self) -> None:
        view = ledger_view(
            [
                record(),
                record(user="petr", kind="ESA", cached_prompt_tokens=None),
                record(model="mystery-model", at=NOW - timedelta(days=1)),
            ]
        )
        assert view.total.calls == 3
        assert view.total.inexact_calls == 1
        assert view.unpriced_models == ["mystery-model"]
        assert [b.key for b in view.by_user][0] in ("jana", "petr")
        assert {b.key for b in view.by_kind} == {"NACE", "ESA"}
        assert [b.key for b in view.by_day] == ["2026-09-24", "2026-09-23"]
        assert view.calls[0].record.at == NOW  # newest first
        assert view.users == ["jana", "petr"]

    def test_cost_matches_the_workbook_pricing(self) -> None:
        view = ledger_view([record()])
        expected = (500 * 0.20 + 500 * 0.02 + 100 * 1.20) / 1_000_000
        assert view.total.cost == pytest.approx(expected)
        assert view.total.cost_per_call == pytest.approx(expected)

    def test_filters(self) -> None:
        rows = [record(), record(user="petr"), record(at=NOW - timedelta(days=3))]
        assert ledger_view(rows, LedgerFilter(user="petr")).matched == 1
        assert ledger_view(rows, LedgerFilter(since=date(2026, 9, 24))).matched == 2
        assert ledger_view(rows, LedgerFilter(until=date(2026, 9, 22))).matched == 1
        assert ledger_view(rows, LedgerFilter(model="nope")).matched == 0
        assert ledger_view(rows, LedgerFilter(kind="NACE")).matched == 3
        assert not LedgerFilter().active and LedgerFilter(user="x").active

    def test_the_row_list_is_capped_but_the_totals_are_not(self) -> None:
        view = ledger_view([record(at=NOW + timedelta(seconds=i)) for i in range(7)], limit=5)
        assert len(view.calls) == 5 and view.matched == 7 and view.truncated

    def test_parse_day(self) -> None:
        assert parse_day("2026-09-24") == date(2026, 9, 24)
        assert parse_day("") is None and parse_day("junk") is None


def a_report(note: str, user: str = "jana", when: datetime = NOW):
    from core.suggest import SuggestionRequest

    provider = StubLlmProvider({"NACE": answer("64"), "ESA": answer("2002703")})
    suggestion = make_service(build_codebooks(), provider=provider).suggest(
        SuggestionRequest(name="Nordkap Funding B.V.")
    )
    return build_report(suggestion, note=note, user=user, now=when)


class TestReportsView:
    def test_newest_first_with_counts(self) -> None:
        older = a_report("první", when=NOW - timedelta(hours=1))
        newer = a_report("druhá", user="petr")
        view = reports_view([older, newer])
        assert [r.note for r in view.reports] == ["druhá", "první"]
        assert view.by_user == [("jana", 1), ("petr", 1)]
        assert view.total == 2

    def test_text_and_user_filters(self) -> None:
        items = [a_report("kaptivní společnost"), a_report("banka", user="petr")]
        assert reports_view(items, ReportFilter(text="KAPTIVNÍ")).matched == 1
        assert reports_view(items, ReportFilter(text="nordkap")).matched == 2
        assert reports_view(items, ReportFilter(user="petr")).matched == 1
        assert reports_view(items, ReportFilter(since=date(2026, 9, 25))).matched == 0


def wire(tmp_path: Path, **settings) -> tuple[TestClient, Database]:
    url = f"sqlite:///{(tmp_path / 'central.sqlite3').as_posix()}"
    config = Settings(llm_api_key=None, llm_cache_path=None, database_url=url, **settings)
    api._state["settings"] = config
    provider = StubLlmProvider({"NACE": answer("64"), "ESA": answer("2002703")})
    api._state["service"] = make_service(build_codebooks(), provider=provider)
    database = Database(url)
    api._state["reports"] = DatabaseReportStore(database)
    return TestClient(api.app), database


def seed(database: Database) -> None:
    ledger = DatabaseLedger(database)
    ledger.record(
        model="gpt-5.6-luna",
        kind="NACE",
        prompt_tokens=3000,
        completion_tokens=120,
        cached_prompt_tokens=1000,
        user="jana",
    )
    ledger.record(
        model="gpt-5.6-luna",
        kind="ESA",
        prompt_tokens=1500,
        completion_tokens=90,
        cached_prompt_tokens=0,
        user="petr",
    )
    DatabaseReportStore(database).save(a_report("Spíš kaptivní společnost.", user="jana"))
    DatabaseAuditSink(database).write(
        {"event": "lookup", "identifier": "X", "user": "jana", "outcome": "found"}
    )


@pytest.fixture
def client(tmp_path: Path) -> Iterator[tuple[TestClient, Database]]:
    try:
        yield wire(tmp_path, admin_password_hash=ADMIN_HASH, session_secret="s3cret")
    finally:
        api._state.clear()


def dev_sign_in(client: TestClient, password: str = ADMIN_PASSWORD, **form):
    return client.post("/admin/login", data={"password": password, **form}, follow_redirects=False)


class TestGate:
    def test_without_a_hash_the_page_does_not_exist(self, tmp_path: Path) -> None:
        try:
            test_client, _ = wire(tmp_path)
            assert test_client.get("/admin").status_code == 404
            assert test_client.get("/admin/login").status_code == 404
            assert test_client.get("/admin/usage.xlsx").status_code == 404
            assert "Pro vývojáře" not in test_client.get("/").text
        finally:
            api._state.clear()

    def test_a_stranger_is_sent_to_the_developer_sign_in(self, client) -> None:
        test_client, _ = client
        response = test_client.get("/admin?user=jana", follow_redirects=False)
        assert response.status_code == 303
        assert response.headers["location"].startswith("/admin/login?next=")
        assert "user%3Djana" in response.headers["location"]
        assert test_client.get("/admin/usage.xlsx", follow_redirects=False).status_code == 303

    def test_the_link_is_on_the_page(self, client) -> None:
        test_client, _ = client
        page = test_client.get("/").text
        assert 'href="/admin"' in page and "Pro vývojáře" in page

    def test_a_wrong_password_is_refused(self, client) -> None:
        test_client, _ = client
        response = dev_sign_in(test_client, "nope")
        assert response.status_code == 401
        assert "Nesprávné heslo" in response.text
        assert "nace_esa_admin" not in response.cookies

    def test_the_right_password_opens_the_page(self, client) -> None:
        test_client, _ = client
        response = dev_sign_in(test_client, next="/admin?since=2026-09-01")
        assert response.status_code == 303
        assert response.headers["location"] == "/admin?since=2026-09-01"
        assert test_client.get("/admin").status_code == 200

    def test_next_never_leaves_the_developer_page(self, client) -> None:
        test_client, _ = client
        assert dev_sign_in(test_client, next="/").headers["location"] == "/admin"
        assert dev_sign_in(test_client, next="https://evil").headers["location"] == "/admin"

    def test_logout_closes_it(self, client) -> None:
        test_client, _ = client
        dev_sign_in(test_client)
        assert test_client.post("/admin/logout", follow_redirects=False).status_code == 303
        assert test_client.get("/admin", follow_redirects=False).status_code == 303

    def test_a_hash_without_a_secret_lets_nobody_in(self, tmp_path: Path) -> None:
        try:
            test_client, _ = wire(tmp_path, admin_password_hash=ADMIN_HASH)
            response = dev_sign_in(test_client)
            assert response.status_code == 503
            assert "SESSION_SECRET" in response.text
        finally:
            api._state.clear()

    def test_the_mo_sign_in_comes_first_when_it_is_on(self, tmp_path: Path) -> None:
        try:
            test_client, _ = wire(
                tmp_path,
                admin_password_hash=ADMIN_HASH,
                app_password_hash=MO_HASH,
                session_secret="s3cret",
            )
            response = test_client.get("/admin/login", follow_redirects=False)
            assert response.status_code == 303 and response.headers["location"].startswith("/login")
            test_client.post(
                "/login", data={"name": "jana", "password": MO_PASSWORD}, follow_redirects=False
            )
            assert test_client.get("/admin/login").status_code == 200
            dev_sign_in(test_client)
            page = test_client.get("/admin").text
            assert "jana" in page  # the developer session carries the MO name
        finally:
            api._state.clear()


class TestPage:
    def test_the_ledger_is_priced_and_grouped(self, client) -> None:
        test_client, database = client
        seed(database)
        dev_sign_in(test_client)
        page = test_client.get("/admin").text
        assert "Ledger nákladů" in page and "Hlášení uživatelů" in page
        assert "jana" in page and "petr" in page
        assert "gpt-5.6-luna" in page
        expected = (
            (2000 * 0.20 + 1000 * 0.02 + 120 * 1.20) + (1500 * 0.20 + 90 * 1.20)
        ) / 1_000_000
        assert f"{expected:.4f}" in page
        assert "Podle uživatele" in page and "Podle dne" in page

    def test_the_filter_narrows_the_ledger(self, client) -> None:
        test_client, database = client
        seed(database)
        dev_sign_in(test_client)
        page = test_client.get("/admin?user=petr").text
        assert "Zrušit filtr" in page
        assert 'value="petr" selected' in page
        calls = page[page.find("<h2>Volání") : page.find('id="complaints"')]
        assert "<td>petr</td>" in calls and "<td>jana</td>" not in calls

    def test_the_complaints_are_listed_with_the_result(self, client) -> None:
        test_client, database = client
        seed(database)
        dev_sign_in(test_client)
        page = test_client.get("/admin").text
        assert "Spíš kaptivní společnost." in page
        assert "Nordkap Funding B.V." in page
        assert "<code>2002703</code>" in page and "<code>64</code>" in page
        assert "error_report" in page  # the JSON detail

    def test_the_complaint_filter(self, client) -> None:
        test_client, database = client
        seed(database)
        dev_sign_in(test_client)
        assert "Spíš kaptivní" in test_client.get("/admin?q=kaptivn").text
        assert "Spíš kaptivní" not in test_client.get("/admin?q=banka").text
        assert "Žádné hlášení pro tento filtr" in test_client.get("/admin?ruser=nobody").text

    def test_an_empty_store_is_an_empty_page_not_an_error(self, client) -> None:
        test_client, _ = client
        dev_sign_in(test_client)
        page = test_client.get("/admin").text
        assert "Žádné volání" in page and "Žádné hlášení" in page

    def test_the_downloads(self, client) -> None:
        test_client, database = client
        seed(database)
        dev_sign_in(test_client)
        usage = test_client.get("/admin/usage.xlsx")
        assert usage.status_code == 200
        book = load_workbook(io.BytesIO(usage.content))
        assert {"Summary", "Calls", "Prices"} <= set(book.sheetnames)
        reports = test_client.get("/admin/reports.xlsx")
        assert reports.status_code == 200
        text = " ".join(
            str(c.value)
            for ws in load_workbook(io.BytesIO(reports.content))
            for row in ws.iter_rows()
            for c in row
        )
        assert "kaptivní" in text and "2002703" in text

    def test_the_audit_log_does_not_gain_a_row_for_reading(self, client) -> None:
        test_client, database = client
        seed(database)
        dev_sign_in(test_client)
        test_client.get("/admin")
        assert database.query("SELECT COUNT(*) FROM audit_events") == [(1,)]
