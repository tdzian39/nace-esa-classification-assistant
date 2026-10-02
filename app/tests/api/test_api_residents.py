"""A Czech issuer through the page, the JSON API and the download: the IČO in the name field,
RES's codes and name on screen, the audit line carrying the IČO - and nothing more.

The real GLEIF and ARES adapters run over mock transports (live payloads of 2 Oct 2026); the
model is off, which is exactly when RES's codes must stand on their own.
"""

from __future__ import annotations

import io
from collections.abc import Iterator

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api import main as api
from config.settings import Settings
from core.classify.llm import LlmClassifier
from core.classify.provider import NullLlmProvider
from core.sources.ares import AresSource
from core.sources.gleif import GleifSource
from core.sources.identity import IssuerIdentifier
from core.sources.web import StaticSearchProvider, WebEvidenceGatherer
from core.suggest import SuggestionService
from tests.classify.conftest import build_resident_codebooks
from tests.sources.conftest import GLEIF_CEZ, RES_CEZ, ares_client, gleif_client

CEZ_ICO = "45274649"
CEZ_ISIN = "CZ0005112300"


@pytest.fixture
def client() -> Iterator[TestClient]:
    settings = Settings(
        web_min_interval_seconds=0.0,
        llm_cache_path=None,
        llm_api_key=None,
        wikimedia_enabled=False,
        gleif_min_interval_seconds=0.0,
        gleif_fetch_parents=False,
        openfigi_enabled=False,
        ares_enabled=True,
        ares_min_interval_seconds=0.0,
    )
    identifier = IssuerIdentifier(
        settings,
        gleif=GleifSource(
            settings,
            client=gleif_client(by_isin={CEZ_ISIN: GLEIF_CEZ}, by_ico={CEZ_ICO: [GLEIF_CEZ]}),
        ),
        ares=AresSource(settings, client=ares_client({CEZ_ICO: RES_CEZ})),
    )
    gatherer = WebEvidenceGatherer(
        settings,
        provider=StaticSearchProvider({}),
        client=httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(404))),
        sleep=lambda _: None,
    )
    api._state["settings"] = Settings(llm_api_key=None, llm_cache_path=None)
    api._state["service"] = SuggestionService(
        build_resident_codebooks(),
        gatherer=gatherer,
        classifier=LlmClassifier(NullLlmProvider()),
        identifier=identifier,
    )
    try:
        yield TestClient(api.app)
    finally:
        api._state.clear()


class TestPage:
    def test_an_ico_in_the_name_field_gives_res_s_codes(self, client: TestClient) -> None:
        page = client.post("/suggest", data={"name": CEZ_ICO}).text
        assert "ČEZ, a. s." in page
        assert f"rezident ČR · IČO {CEZ_ICO}" in page
        assert "číselník BA0036 · jen elementární kódy · rezident" in page
        assert "převzato z RES (ARES), stav k 2026-09-04" in page
        assert "z RES (ARES) · ověřte" in page and "podle pravidel · ověřte" not in page
        assert "1100100" in page

    def test_the_resident_warning_is_gone(self, client: TestClient) -> None:
        page = client.post("/suggest", data={"isin": CEZ_ISIN}).text
        assert "nabízí jen nerezidentské kódy" not in page
        assert "kódy z RES převzít nešlo" not in page

    def test_the_download_and_the_report_forms_keep_the_ico(self, client: TestClient) -> None:
        page = client.post("/suggest", data={"name": CEZ_ICO}).text
        assert page.count(f'type="hidden" name="name" value="{CEZ_ICO}"') == 2

    def test_the_form_says_an_ico_may_be_typed(self, client: TestClient) -> None:
        page = client.get("/").text
        assert 'Název emitenta <span class="note">nebo IČO</span>' in page
        assert "Zadejte ISIN, název nebo IČO emitenta, nebo popis činnosti." in page


class TestJson:
    @pytest.mark.parametrize("body", [{"name": CEZ_ICO}, {"ico": CEZ_ICO}])
    def test_the_identity_names_the_res_record(self, client: TestClient, body: dict) -> None:
        data = client.post("/api/suggest", json=body).json()
        identity = data["identity"]
        assert identity["resident"] is True and identity["ico"] == CEZ_ICO
        assert identity["res"] == {
            "name": "ČEZ, a. s.",
            "legal_form": "121",
            "sector": "11001",
            "nace": "35110",
            "nace_2008": "35110",
            "as_of": "2026-09-04",
        }
        assert data["esa"]["proposal"]["code"] == "1100100"
        assert data["nace"]["proposal"]["code"] == "35"
        assert data["row"]["source"] == "GLEIF+RES+WEB"

    def test_the_row_echoes_the_ico_where_it_was_typed(self, client: TestClient) -> None:
        row = client.post("/api/suggest", json={"name": CEZ_ICO}).json()["row"]
        assert row["IN_name"] == CEZ_ICO and row["issuer_name"] == "ČEZ, a. s."


def test_the_audit_line_carries_the_ico_and_counts_res_as_found(
    client: TestClient, caplog: pytest.LogCaptureFixture
) -> None:
    with caplog.at_level("INFO", logger="core.audit"):
        client.post("/suggest", data={"name": CEZ_ICO})
    (line,) = [r.message for r in caplog.records if r.message.startswith("lookup identifier=")]
    assert f"ico={CEZ_ICO}" in line and "outcome=found" in line
    assert "sources=GLEIF+RES+WEB" in line
    assert "11001" not in line and "35110" not in line, "never what RES said"


def test_the_download_has_res_s_codes_and_the_ico_on_the_run_sheet(client: TestClient) -> None:
    response = client.get("/suggest.xlsx", params={"name": CEZ_ICO})
    workbook = load_workbook(io.BytesIO(response.content))
    main, run = workbook.worksheets[0], workbook.worksheets[1]
    rows = list(main.iter_rows(values_only=True))
    assert rows[1][0] == "ČEZ, a. s."
    assert "35" in rows[1] and "1100100" in rows[1]
    pairs = {row[0]: row[1] for row in run.iter_rows(values_only=True) if row and row[0]}
    assert pairs["IČO"] == CEZ_ICO
