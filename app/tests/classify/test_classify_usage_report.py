"""``--usage-xlsx``: the usage ledger as a workbook a person can read, with honest costs."""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from config.settings import Settings
from core.classify import __main__ as cli
from core.classify.budget import SqliteLedger, UsageRecord
from core.classify.usage_report import (
    CALL_COLUMNS,
    CALLS_SHEET,
    PRICES,
    PRICES_SHEET,
    SUMMARY_SHEET,
    price_for,
    write_usage_workbook,
)
from core.codebooks.models import CodebookSet

T0 = datetime(2026, 9, 23, 7, 53, 17, tzinfo=UTC)

#: 3,000 input and 100 output tokens at gpt-5.6-luna's $0.20 / $1.20 per million.
ONE_LUNA_CALL = 3_000 * 0.20 / 1_000_000 + 100 * 1.20 / 1_000_000


def record(
    model: str = "gpt-5.6-luna",
    kind: str = "NACE",
    prompt_tokens: int = 3_000,
    completion_tokens: int = 100,
    minutes: int = 0,
) -> UsageRecord:
    return UsageRecord(
        at=T0 + timedelta(minutes=minutes),
        model=model,
        kind=kind,
        prompt_tokens=prompt_tokens,
        completion_tokens=completion_tokens,
    )


def sheet_text(path: Path, sheet: str) -> list[object]:
    return [
        cell.value for row in load_workbook(path)[sheet].iter_rows() for cell in row if cell.value
    ]


@pytest.fixture
def out(tmp_path: Path) -> Path:
    return tmp_path / "reports" / "usage.xlsx"  # the folder does not exist yet


class TestPrices:
    def test_a_known_model_has_its_price(self) -> None:
        assert price_for("gpt-5.6-luna") == PRICES["gpt-5.6-luna"]

    def test_a_dated_snapshot_is_priced_as_its_model(self) -> None:
        assert price_for("gpt-6-luna-2026-08-01") == PRICES["gpt-6-luna"]

    def test_an_unknown_model_has_no_price_rather_than_a_guess(self) -> None:
        assert price_for("gpt-4o-mini") is None

    def test_a_longer_name_is_not_a_snapshot(self) -> None:
        assert price_for("gpt-6-lunar") is None


class TestLedgerRecords:
    def test_calls_come_back_oldest_first_with_their_fields(self, tmp_path: Path) -> None:
        ledger = SqliteLedger(tmp_path / "usage.sqlite3")
        ledger.record(model="gpt-6-luna", kind="NACE", prompt_tokens=3630, completion_tokens=74)
        ledger.record(model="gpt-6-luna", kind="ESA", prompt_tokens=1278, completion_tokens=99)

        records = ledger.records()

        assert [(r.kind, r.prompt_tokens, r.completion_tokens) for r in records] == [
            ("NACE", 3630, 74),
            ("ESA", 1278, 99),
        ]
        assert all(r.at.tzinfo is not None for r in records)
        assert records[0].at <= records[1].at

    def test_an_unusable_ledger_has_no_records(self, tmp_path: Path) -> None:
        blocker = tmp_path / "blocker"
        blocker.write_text("not a directory", encoding="utf-8")
        assert SqliteLedger(blocker / "nested" / "usage.sqlite3").records() == []


class TestWorkbook:
    def test_it_has_a_summary_the_calls_and_the_prices(self, out: Path) -> None:
        write_usage_workbook([record()], out)
        assert load_workbook(out).sheetnames == [SUMMARY_SHEET, CALLS_SHEET, PRICES_SHEET]

    def test_one_row_per_call_with_its_cost(self, out: Path) -> None:
        write_usage_workbook([record()], out)
        sheet = load_workbook(out)[CALLS_SHEET]

        assert tuple(cell.value for cell in sheet[1]) == CALL_COLUMNS
        row = {name: cell.value for name, cell in zip(CALL_COLUMNS, sheet[2], strict=True)}
        assert row["Time (UTC)"] == datetime(2026, 9, 23, 7, 53, 17)  # naive UTC, as elsewhere
        assert (row["Model"], row["Codebook"], row["Total tokens"]) == (
            "gpt-5.6-luna",
            "NACE",
            3100,
        )
        assert row["Cost (USD)"] == pytest.approx(ONE_LUNA_CALL)

    def test_calls_are_listed_in_time_order(self, out: Path) -> None:
        write_usage_workbook([record(kind="ESA", minutes=5), record(kind="NACE")], out)
        sheet = load_workbook(out)[CALLS_SHEET]
        assert [sheet.cell(row=r, column=4).value for r in (2, 3)] == ["NACE", "ESA"]

    def test_totals_by_model_codebook_and_day(self, out: Path) -> None:
        records = [
            record(model="gpt-6-luna", kind="NACE"),
            record(model="gpt-6-luna", kind="ESA", minutes=1),
            record(kind="ESA", minutes=2),
        ]
        report = write_usage_workbook(records, out)

        text = sheet_text(out, SUMMARY_SHEET)
        assert {"By model", "By codebook", "By day", "gpt-6-luna", "ESA", "2026-09-23"} <= set(text)
        two_cheaper_calls = 2 * (3_000 * 0.10 + 100 * 0.50) / 1_000_000
        assert report.calls == 3
        assert report.cost == pytest.approx(two_cheaper_calls + ONE_LUNA_CALL)

    def test_an_unknown_model_gets_empty_costs_and_is_named(self, out: Path) -> None:
        report = write_usage_workbook([record(), record(model="mystery-model", minutes=1)], out)

        assert load_workbook(out)[CALLS_SHEET].cell(row=3, column=10).value is None
        assert report.unpriced_models == ("mystery-model",)
        assert report.cost == pytest.approx(ONE_LUNA_CALL)  # the priced call still counts
        assert any(
            "No price known for mystery-model" in str(v) for v in sheet_text(out, SUMMARY_SHEET)
        )

    def test_the_summary_says_what_the_total_is_not(self, out: Path) -> None:
        """Production keeps no ledger and cached input is not recorded: the total is not the bill."""
        write_usage_workbook([record()], out)
        text = " ".join(str(value) for value in sheet_text(out, SUMMARY_SHEET))
        assert "Production" in text and "NOT here" in text
        assert "upper bound" in text

    def test_the_prices_sheet_lists_every_price_and_its_source(self, out: Path) -> None:
        write_usage_workbook([record()], out)
        sheet = load_workbook(out)[PRICES_SHEET]
        models = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]
        assert models == sorted(PRICES)
        assert str(sheet.cell(row=2, column=5).value).startswith("https://")

    def test_an_empty_ledger_still_makes_a_valid_workbook(self, out: Path) -> None:
        report = write_usage_workbook([], out)
        assert report.calls == 0 and report.cost == 0
        assert load_workbook(out)[CALLS_SHEET].max_row == 1


class TestCommand:
    """``python -m core.classify --usage-xlsx [PATH]``."""

    @staticmethod
    def _no_codebooks(settings: object) -> CodebookSet:
        raise AssertionError("the usage export needs no codebooks")

    @pytest.fixture
    def ledger_path(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
        path = tmp_path / "llm_usage.sqlite3"
        monkeypatch.setattr(
            cli, "get_settings", lambda: Settings(_env_file=None, llm_usage_path=str(path))
        )
        monkeypatch.setattr(cli, "_load", self._no_codebooks)
        return path

    def test_it_writes_the_workbook_next_to_the_ledger(
        self, ledger_path: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        SqliteLedger(ledger_path).record(
            model="gpt-5.6-luna", kind="NACE", prompt_tokens=3_000, completion_tokens=100
        )
        assert cli.main(["--usage-xlsx"]) == cli.EXIT_OK
        assert ledger_path.with_suffix(".xlsx").exists()
        output = capsys.readouterr().out
        assert "1 call(s)" in output and "upper bound" in output

    def test_a_path_given_is_honoured(self, ledger_path: Path, tmp_path: Path) -> None:
        SqliteLedger(ledger_path).record(
            model="gpt-6-luna", kind="ESA", prompt_tokens=1_000, completion_tokens=50
        )
        target = tmp_path / "elsewhere" / "report.xlsx"
        assert cli.main(["--usage-xlsx", str(target)]) == cli.EXIT_OK
        assert target.exists()

    def test_an_unpriced_model_is_named(
        self, ledger_path: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        SqliteLedger(ledger_path).record(
            model="mystery-model", kind="NACE", prompt_tokens=1, completion_tokens=1
        )
        assert cli.main(["--usage-xlsx"]) == cli.EXIT_OK
        assert "no price for mystery-model" in capsys.readouterr().out

    def test_nothing_recorded_yet_writes_nothing(
        self, ledger_path: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        assert cli.main(["--usage-xlsx"]) == cli.EXIT_OK
        assert "nothing recorded yet" in capsys.readouterr().out
        assert not ledger_path.exists()
        assert not ledger_path.with_suffix(".xlsx").exists()

    def test_recording_switched_off_is_an_error(
        self, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
    ) -> None:
        monkeypatch.setattr(
            cli, "get_settings", lambda: Settings(_env_file=None, llm_usage_path="")
        )
        monkeypatch.setattr(cli, "_load", self._no_codebooks)
        assert cli.main(["--usage-xlsx"]) == cli.EXIT_LOAD_FAILED
        assert "LLM_USAGE_PATH is empty" in capsys.readouterr().err

    def test_an_unwritable_target_is_an_error_not_a_crash(
        self, ledger_path: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        """What a workbook still open in Excel looks like on Windows: the save is refused."""
        SqliteLedger(ledger_path).record(
            model="gpt-5.6-luna", kind="NACE", prompt_tokens=1, completion_tokens=1
        )
        folder = tmp_path / "a-folder.xlsx"
        folder.mkdir()
        assert cli.main(["--usage-xlsx", str(folder)]) == cli.EXIT_LOAD_FAILED
        assert "could not write" in capsys.readouterr().err
