"""The result workbook: cell types, number formats and the Run sheet.

Every test reads the written file back with openpyxl, so what is asserted is what Excel
would actually show.
"""

from __future__ import annotations

from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import load_workbook

from core.export.columns import SUBJECT_COLUMNS, record_row
from core.export.xlsx import MAX_CELL_LENGTH, RUN_SHEET, SUBJECTS_SHEET, write_workbook
from tests.batch.conftest import record


def write(tmp_path: Path, columns, rows, **kwargs) -> Path:
    return write_workbook(tmp_path / "out.xlsx", columns, rows, **kwargs)


class TestStructure:
    def test_header_and_one_row_per_subject(self, tmp_path: Path) -> None:
        rows = [record_row(record(), status="found"), record_row(None, status="not_found")]
        sheet = load_workbook(write(tmp_path, SUBJECT_COLUMNS, rows))[SUBJECTS_SHEET]

        assert sheet.max_row == 3  # header + two subjects
        assert sheet.cell(row=1, column=1).value == "status"

    def test_header_is_frozen_and_filterable(self, tmp_path: Path) -> None:
        sheet = load_workbook(write(tmp_path, SUBJECT_COLUMNS, [record_row(record())]))[
            SUBJECTS_SHEET
        ]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None

    def test_missing_keys_become_empty_cells(self, tmp_path: Path) -> None:
        sheet = load_workbook(write(tmp_path, ("a", "b"), [{"a": 1}]))[SUBJECTS_SHEET]
        assert sheet.cell(row=2, column=1).value == 1
        assert sheet.cell(row=2, column=2).value is None

    def test_no_run_sheet_unless_metadata_is_given(self, tmp_path: Path) -> None:
        book = load_workbook(write(tmp_path, SUBJECT_COLUMNS, [record_row(record())]))
        assert book.sheetnames == [SUBJECTS_SHEET]


class TestCellTypes:
    def _cell(self, tmp_path: Path, column: str):
        rows = [record_row(record(), status="found")]
        sheet = load_workbook(write(tmp_path, SUBJECT_COLUMNS, rows))[SUBJECTS_SHEET]
        index = SUBJECT_COLUMNS.index(column) + 1
        return sheet.cell(row=2, column=index)

    def test_ico_stays_text_with_its_leading_zeros(self, tmp_path: Path) -> None:
        """The whole point of Tool 2 is fixing lost leading zeros; the export must not lose them."""
        rows = [record_row(record(ico="00177041"), status="found")]
        sheet = load_workbook(write(tmp_path, SUBJECT_COLUMNS, rows))[SUBJECTS_SHEET]
        cell = sheet.cell(row=2, column=SUBJECT_COLUMNS.index("ico") + 1)

        assert cell.value == "00177041"
        assert isinstance(cell.value, str)
        assert cell.number_format == "@"

    def test_nace_codes_stay_text(self, tmp_path: Path) -> None:
        cell = self._cell(tmp_path, "RES_nace_rev2_main")
        assert cell.value == "64190"
        assert cell.number_format == "@"

    def test_dates_are_real_dates(self, tmp_path: Path) -> None:
        cell = self._cell(tmp_path, "RES_founded_on")
        assert cell.value == datetime(1993, 6, 25)
        assert cell.number_format == "yyyy-mm-dd"

    def test_booleans_stay_boolean(self, tmp_path: Path) -> None:
        assert self._cell(tmp_path, "nace_mismatch").value is False

    def test_unknown_mismatch_is_empty_not_false(self, tmp_path: Path) -> None:
        """None means "cannot tell"; writing FALSE would assert the revisions agree."""
        subject = record()
        stripped = subject.res.__class__(
            ico=subject.res.ico, provenance=subject.res.provenance, name=subject.res.name
        )
        rows = [record_row(subject.__class__(ico=subject.ico, res=stripped), status="found")]
        sheet = load_workbook(write(tmp_path, SUBJECT_COLUMNS, rows))[SUBJECTS_SHEET]
        assert sheet.cell(row=2, column=SUBJECT_COLUMNS.index("nace_mismatch") + 1).value is None


class TestListJoining:
    def _value(self, tmp_path: Path, column: str) -> str:
        rows = [record_row(record(), status="found")]
        sheet = load_workbook(write(tmp_path, SUBJECT_COLUMNS, rows))[SUBJECTS_SHEET]
        return sheet.cell(row=2, column=SUBJECT_COLUMNS.index(column) + 1).value

    def test_codes_are_separated_by_semicolons(self, tmp_path: Path) -> None:
        assert self._value(tmp_path, "RES_nace_rev2_other") == "66190"

    def test_long_texts_are_separated_by_newlines(self, tmp_path: Path) -> None:
        value = self._value(tmp_path, "OR_predmet_podnikani")
        assert value == "bankovní obchody\npronájem nemovitostí"

    def test_wrapped_columns_wrap(self, tmp_path: Path) -> None:
        rows = [record_row(record(), status="found")]
        sheet = load_workbook(write(tmp_path, SUBJECT_COLUMNS, rows))[SUBJECTS_SHEET]
        cell = sheet.cell(row=2, column=SUBJECT_COLUMNS.index("OR_predmet_podnikani") + 1)
        assert cell.alignment.wrap_text is True

    def test_an_empty_list_is_an_empty_cell(self, tmp_path: Path) -> None:
        sheet = load_workbook(write(tmp_path, ("a",), [{"a": ()}]))[SUBJECTS_SHEET]
        assert sheet.cell(row=2, column=1).value is None


class TestHostileText:
    def test_illegal_control_characters_are_stripped(self, tmp_path: Path) -> None:
        """openpyxl refuses these outright; a register text must not crash the export."""
        path = write(tmp_path, ("a",), [{"a": "bankovní\x0bobchody"}])
        assert load_workbook(path)[SUBJECTS_SHEET].cell(row=2, column=1).value == (
            "bankovníobchody"
        )

    def test_over_long_text_is_truncated_not_rejected(self, tmp_path: Path) -> None:
        path = write(tmp_path, ("a",), [{"a": "x" * (MAX_CELL_LENGTH + 500)}])
        value = load_workbook(path)[SUBJECTS_SHEET].cell(row=2, column=1).value

        assert len(value) <= MAX_CELL_LENGTH
        assert value.endswith("[…]")


class TestRunSheet:
    def test_metadata_is_written_as_key_value_pairs(self, tmp_path: Path) -> None:
        metadata = {
            "input file": "klienti.xlsx",
            "rows": 2,
            "requested by": "tester",
            "sources": ("DWS", "ARES_LIVE"),
            "codebook version": "cb-0123456789abcdef",
        }
        book = load_workbook(write(tmp_path, SUBJECT_COLUMNS, [], run_metadata=metadata))
        sheet = book[RUN_SHEET]
        pairs = {
            sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
            for r in range(2, sheet.max_row + 1)
        }

        assert pairs["input file"] == "klienti.xlsx"
        assert pairs["sources"] == "DWS, ARES_LIVE"
        assert pairs["codebook version"] == "cb-0123456789abcdef"

    def test_datetimes_survive(self, tmp_path: Path) -> None:
        metadata = {
            "started at (UTC)": datetime(2026, 9, 22, 12, 0, tzinfo=UTC).replace(tzinfo=None)
        }
        book = load_workbook(write(tmp_path, SUBJECT_COLUMNS, [], run_metadata=metadata))
        assert book[RUN_SHEET].cell(row=2, column=2).value == datetime(2026, 9, 22, 12, 0)


def test_parent_directories_are_created(tmp_path: Path) -> None:
    path = write_workbook(tmp_path / "new" / "dir" / "out.xlsx", ("a",), [{"a": 1}])
    assert path.is_file()


def test_date_column_written_from_a_plain_date(tmp_path: Path) -> None:
    path = write_workbook(
        tmp_path / "d.xlsx", ("RES_founded_on",), [{"RES_founded_on": date(2020, 1, 2)}]
    )
    assert load_workbook(path)[SUBJECTS_SHEET].cell(row=2, column=1).value == datetime(2020, 1, 2)
