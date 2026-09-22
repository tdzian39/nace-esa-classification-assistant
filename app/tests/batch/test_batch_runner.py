"""End to end: a messy sheet in, one result row per input row out."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.batch.runner import (
    IN_IDENTIFIER_COLUMN,
    IN_NOTE_COLUMN,
    IN_ROW_COLUMN,
    build_columns,
    default_output_path,
    resolve_rows,
    run_batch,
)
from core.export.columns import SUBJECT_COLUMNS, header_label
from core.export.xlsx import RUN_SHEET, SUBJECTS_SHEET
from core.sources.base import SubjectCandidate
from tests.batch.conftest import (
    BAD_CHECKSUM,
    RAIFFEISENBANK,
    SKODA,
    ScriptedSource,
    make_resolver,
    record,
    write_sheet,
)


def run(input_path: Path, source: ScriptedSource, **kwargs) -> tuple:
    """Run a batch and return ``(report, subjects sheet)``."""
    resolver = make_resolver(source)
    report = run_batch(
        input_path,
        kwargs.pop("output_path", None),
        resolver=resolver,
        user="tester",
        codebook_version="cb-0123456789abcdef",
        **kwargs,
    )
    return report, load_workbook(report.output_path)[SUBJECTS_SHEET]


def header_index(sheet, name: str) -> int:
    """Index of a column by its key; the sheet stores the human label (``ico`` -> ``IČO``)."""
    wanted = header_label(name)
    for index in range(1, sheet.max_column + 1):
        if sheet.cell(row=1, column=index).value == wanted:
            return index
    raise AssertionError(f"column {name!r} (header {wanted!r}) not in the sheet")


def column_values(sheet, name: str) -> list:
    index = header_index(sheet, name)
    return [sheet.cell(row=r, column=index).value for r in range(2, sheet.max_row + 1)]


class TestRowsOut:
    def test_one_row_per_input_row_in_order(self, messy_input: Path) -> None:
        source = ScriptedSource({SKODA: record(SKODA), RAIFFEISENBANK: record()}, default=None)
        report, sheet = run(messy_input, source)

        assert report.row_count == 3
        assert sheet.max_row == 4  # header + three rows
        assert column_values(sheet, IN_ROW_COLUMN) == [3, 4, 6]

    def test_unresolved_rows_are_still_written(self, messy_input: Path) -> None:
        """A reviewer must see the empty lines, not find them missing from the sheet."""
        source = ScriptedSource({SKODA: record(SKODA)}, default=None)
        _, sheet = run(messy_input, source)
        assert column_values(sheet, "status") == ["found", "not_found", "not_found"]

    def test_lost_leading_zeros_are_repaired_on_the_way_through(self, messy_input: Path) -> None:
        """Input said 177041; the row must come back as Škoda Auto with a full 8-digit IČO."""
        source = ScriptedSource({SKODA: record(SKODA, name="Škoda Auto a.s.")}, default=None)
        _, sheet = run(messy_input, source)

        assert source.ico_calls[0] == SKODA
        assert column_values(sheet, "ico")[0] == SKODA
        assert column_values(sheet, "RES_name")[0] == "Škoda Auto a.s."

    def test_the_status_of_a_bad_ico_explains_itself(self, tmp_path: Path) -> None:
        path = write_sheet(
            tmp_path / "bad.xlsx", [("IČO", "Název"), (BAD_CHECKSUM, "Raiffeisenbank a.s.")]
        )
        source = ScriptedSource(default=record())
        _, sheet = run(path, source)

        assert column_values(sheet, "status") == ["invalid_input"]
        assert source.ico_calls == []  # nothing was looked up under a wrong identifier
        assert "checksum" in column_values(sheet, "notes")[0]
        assert "not a valid IČO" in column_values(sheet, IN_NOTE_COLUMN)[0]

    def test_ambiguous_names_list_their_candidates(self, tmp_path: Path) -> None:
        path = write_sheet(tmp_path / "name.xlsx", [("Název",), ("Raiffeisen",)])
        source = ScriptedSource(
            candidates=(
                SubjectCandidate(RAIFFEISENBANK, "Raiffeisenbank a.s.", "ARES_LIVE"),
                SubjectCandidate("26400276", "Raiffeisen stavební spořitelna a.s.", "ARES_LIVE"),
            )
        )
        _, sheet = run(path, source)

        assert column_values(sheet, "status") == ["ambiguous"]
        notes = column_values(sheet, "notes")[0]
        assert "26400276" in notes and RAIFFEISENBANK in notes


class TestInputEcho:
    def test_input_columns_are_copied_with_an_in_prefix(self, messy_input: Path) -> None:
        _, sheet = run(messy_input, ScriptedSource(default=record()))
        assert column_values(sheet, "IN_Poznámka")[0] == "flag A"

    def test_echo_can_be_switched_off(self, messy_input: Path) -> None:
        _, sheet = run(messy_input, ScriptedSource(default=record()), echo_input=False)
        headers = [sheet.cell(row=1, column=i).value for i in range(1, sheet.max_column + 1)]
        assert "IN_Poznámka" not in headers
        assert IN_IDENTIFIER_COLUMN in headers  # the identifier is always kept

    def test_echoed_columns_cannot_collide_with_subject_columns(self, tmp_path: Path) -> None:
        """An input column literally called "status" must not overwrite the result status."""
        path = write_sheet(tmp_path / "collide.xlsx", [("IČO", "status"), (RAIFFEISENBANK, "open")])
        _, sheet = run(path, ScriptedSource(default=record()))

        assert column_values(sheet, "IN_status") == ["open"]
        assert column_values(sheet, "status") == ["found"]

    def test_columns_are_input_echo_then_subject(self, messy_input: Path) -> None:
        from core.batch.reader import read_batch_input

        columns = build_columns(read_batch_input(messy_input), echo_input=True)
        assert columns[:3] == (IN_ROW_COLUMN, IN_IDENTIFIER_COLUMN, IN_NOTE_COLUMN)
        assert columns[-len(SUBJECT_COLUMNS) :] == SUBJECT_COLUMNS


class TestCaching:
    def test_a_repeated_identifier_is_resolved_once(self, tmp_path: Path) -> None:
        """A monthly list names the same client repeatedly; ARES should be asked once."""
        path = write_sheet(
            tmp_path / "dupes.xlsx",
            [("IČO",), (RAIFFEISENBANK,), (RAIFFEISENBANK,), ("  49 240 901  ",)],
        )
        source = ScriptedSource(default=record())
        report, sheet = run(path, source)

        assert report.row_count == 3
        assert len(source.ico_calls) == 1
        assert column_values(sheet, "status") == ["found", "found", "found"]

    def test_the_cache_is_case_insensitive_for_names(self, tmp_path: Path) -> None:
        path = write_sheet(tmp_path / "names.xlsx", [("Název",), ("Raiffeisen",), ("RAIFFEISEN",)])
        source = ScriptedSource()
        resolve_rows(
            __import__("core.batch.reader", fromlist=["x"]).read_batch_input(path).rows,
            make_resolver(source),
        )
        assert len(source.name_calls) == 1


class TestReport:
    def test_counts_every_status(self, messy_input: Path) -> None:
        source = ScriptedSource({SKODA: record(SKODA)}, default=None)
        report, _ = run(messy_input, source)

        assert report.counts == {"found": 1, "not_found": 2}
        assert report.resolved == 1
        assert report.unresolved == 2

    def test_summary_names_the_output(self, messy_input: Path) -> None:
        report, _ = run(messy_input, ScriptedSource(default=record()))
        assert str(report.output_path) in report.summary()

    def test_run_sheet_records_the_provenance_of_the_run(self, messy_input: Path) -> None:
        report, _ = run(messy_input, ScriptedSource(default=record()))
        sheet = load_workbook(report.output_path)[RUN_SHEET]
        pairs = {
            sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
            for r in range(2, sheet.max_row + 1)
        }

        assert pairs["input file"] == "messy.xlsx"
        assert pairs["requested by"] == "tester"
        assert pairs["codebook version"] == "cb-0123456789abcdef"
        assert pairs["rows"] == 3

    def test_missing_codebooks_are_stated_not_left_blank(self, messy_input: Path) -> None:
        resolver = make_resolver(ScriptedSource(default=record()))
        report = run_batch(
            messy_input, None, resolver=resolver, user="tester", codebook_version=None
        )
        assert "not loaded" in report.run_metadata()["codebook version"]


class TestOutputPath:
    def test_default_is_beside_the_input(self, tmp_path: Path) -> None:
        assert default_output_path(tmp_path / "klienti.xlsx").name == "klienti_lookup.xlsx"

    def test_explicit_path_is_honoured(self, messy_input: Path, tmp_path: Path) -> None:
        destination = tmp_path / "sub" / "vysledek.xlsx"
        report, _ = run(messy_input, ScriptedSource(default=record()), output_path=destination)
        assert report.output_path == destination
        assert destination.is_file()

    def test_overwriting_the_input_is_refused(self, messy_input: Path) -> None:
        """Losing the source list to its own result would be unrecoverable."""
        resolver = make_resolver(ScriptedSource(default=record()))
        with pytest.raises(ValueError, match="overwrite the input"):
            run_batch(messy_input, messy_input, resolver=resolver, user="tester")
