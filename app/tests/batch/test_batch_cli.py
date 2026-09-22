"""``python -m core.batch``: argument handling, exit codes and console output."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from config.settings import Settings
from core.batch import __main__ as cli
from core.export.xlsx import SUBJECTS_SHEET
from core.sources.base import SourceUnavailableError
from core.sources.resolver import SubjectResolver
from tests.batch.conftest import RAIFFEISENBANK, ScriptedSource, record, write_sheet


@pytest.fixture
def patched_cli(monkeypatch: pytest.MonkeyPatch) -> None:
    """In-memory settings and no codebook load, so the CLI runs anywhere."""
    monkeypatch.setattr(cli, "get_settings", lambda: Settings(dws_dsn=None, lookup_user="tester"))
    monkeypatch.setattr(cli, "_codebook_version", lambda settings: "cb-0123456789abcdef")


def use_sources(monkeypatch: pytest.MonkeyPatch, *sources: ScriptedSource) -> None:
    monkeypatch.setattr(
        cli,
        "build_default_resolver",
        lambda settings, **kwargs: SubjectResolver(list(sources), user="tester"),
    )


class TestExitCodes:
    def test_all_rows_resolved_exits_zero(
        self,
        patched_cli: None,
        monkeypatch: pytest.MonkeyPatch,
        clean_input: Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        use_sources(monkeypatch, ScriptedSource(default=record()))
        assert cli.main([str(clean_input)]) == cli.EXIT_OK

    def test_unresolved_rows_exit_one(
        self,
        patched_cli: None,
        monkeypatch: pytest.MonkeyPatch,
        clean_input: Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        use_sources(monkeypatch, ScriptedSource(default=None))
        assert cli.main([str(clean_input)]) == cli.EXIT_INCOMPLETE

    def test_unreadable_input_exits_two(
        self,
        patched_cli: None,
        monkeypatch: pytest.MonkeyPatch,
        tmp_path: Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        use_sources(monkeypatch, ScriptedSource(default=record()))
        assert cli.main([str(tmp_path / "missing.xlsx")]) == cli.EXIT_FAILED
        assert "error:" in capsys.readouterr().err

    def test_no_source_enabled_exits_two(
        self,
        patched_cli: None,
        monkeypatch: pytest.MonkeyPatch,
        clean_input: Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        use_sources(monkeypatch)
        assert cli.main([str(clean_input)]) == cli.EXIT_FAILED

    def test_every_row_failing_exits_two(
        self,
        patched_cli: None,
        monkeypatch: pytest.MonkeyPatch,
        clean_input: Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        """A total source outage is a failed run, not a sheet full of honest misses."""
        source = ScriptedSource()
        source.fetch_by_ico = lambda ico: (_ for _ in ()).throw(SourceUnavailableError("down"))
        use_sources(monkeypatch, source)
        assert cli.main([str(clean_input)]) == cli.EXIT_FAILED


class TestOutput:
    def test_default_output_lands_beside_the_input(
        self,
        patched_cli: None,
        monkeypatch: pytest.MonkeyPatch,
        clean_input: Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        use_sources(monkeypatch, ScriptedSource(default=record()))
        cli.main([str(clean_input)])
        assert clean_input.with_name("clean_lookup.xlsx").is_file()

    def test_explicit_output_path(
        self,
        patched_cli: None,
        monkeypatch: pytest.MonkeyPatch,
        clean_input: Path,
        tmp_path: Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        destination = tmp_path / "vysledek.xlsx"
        use_sources(monkeypatch, ScriptedSource(default=record()))
        cli.main([str(clean_input), "-o", str(destination)])
        assert destination.is_file()

    def test_summary_reports_the_counts(
        self,
        patched_cli: None,
        monkeypatch: pytest.MonkeyPatch,
        clean_input: Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        use_sources(monkeypatch, ScriptedSource(default=None))
        cli.main([str(clean_input)])
        out = capsys.readouterr().out
        assert "not_found=2" in out
        assert "need review" in out

    def test_sheet_can_be_chosen(
        self,
        patched_cli: None,
        monkeypatch: pytest.MonkeyPatch,
        tmp_path: Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        path = write_sheet(
            tmp_path / "multi.xlsx",
            [("Pokyny",)],
            title="Pokyny",
            extra={"Data": [("IČO",), (RAIFFEISENBANK,)]},
        )
        source = ScriptedSource(default=record())
        use_sources(monkeypatch, source)
        cli.main([str(path), "--sheet", "Data"])
        assert source.ico_calls == [RAIFFEISENBANK]

    def test_echo_can_be_switched_off(
        self,
        patched_cli: None,
        monkeypatch: pytest.MonkeyPatch,
        clean_input: Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        use_sources(monkeypatch, ScriptedSource(default=record()))
        cli.main([str(clean_input), "--no-echo-input"])
        sheet = load_workbook(clean_input.with_name("clean_lookup.xlsx"))[SUBJECTS_SHEET]
        headers = [sheet.cell(row=1, column=i).value for i in range(1, sheet.max_column + 1)]
        assert "IN_Název klienta" not in headers
