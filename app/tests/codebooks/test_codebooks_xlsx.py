"""Tests of the tolerant openpyxl table reader."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from core.codebooks.errors import CodebookFileError, CodebookSchemaError
from core.codebooks.loaders import CTS_COLUMNS
from core.codebooks.loaders import VALID_ESA_COLUMNS as VALID_COLUMNS
from core.codebooks.xlsx import Table, fold_header, read_table

from .conftest import corrupt_sheet_xml, write_xlsx


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("Kód", "kod"),
        ("KÓD", "kod"),
        (" Kod ", "kod"),
        ("Název", "nazev"),
        ("Kód  NACE", "kod nace"),
        ("VALUE", "value"),
        (None, ""),
        ("Zkrácený\N{NO-BREAK SPACE}text", "zkraceny text"),
    ],
)
def test_fold_header(text: object, expected: str) -> None:
    assert fold_header(text) == expected


def test_reads_plain_table(tmp_path: Path) -> None:
    path = write_xlsx(
        tmp_path / "cts.xlsx",
        ("ID", "VALUE", "DESCRIPTION"),
        [(1001, "S.11001", "Veřejné"), (1002, "S.11002", "Národní")],
    )
    table = read_table(path, CTS_COLUMNS)
    assert isinstance(table, Table)
    assert table.header_row == 1
    assert table.sheet == "Sheet1"
    assert table.path == path
    assert table.columns == {"ID": "ID", "VALUE": "VALUE", "DESCRIPTION": "DESCRIPTION"}
    assert table.rows == (
        {"ID": "1001", "VALUE": "S.11001", "DESCRIPTION": "Veřejné"},
        {"ID": "1002", "VALUE": "S.11002", "DESCRIPTION": "Národní"},
    )
    assert table.row_numbers == (2, 3)
    assert list(table.iter_numbered()) == list(zip(table.row_numbers, table.rows, strict=True))


def test_header_below_title_rows(tmp_path: Path) -> None:
    path = write_xlsx(
        tmp_path / "cts.xlsx",
        ("ID", "VALUE", "DESCRIPTION"),
        [(1001, "S.11001", "x")],
        title_rows=[("Export CTS 2024-06-30",), (), ("Codebook", "BA0036")],
    )
    table = read_table(path, CTS_COLUMNS)
    assert table.header_row == 4
    assert table.row_numbers == (5,)
    assert table.rows[0]["ID"] == "1001"


@pytest.mark.parametrize(
    "headers",
    [
        ("KOD", "Nazev", "Popis"),
        ("kód", "název", "popis"),
        ("Code", "Name", "Description"),
        (" Kód ", "Název\N{NO-BREAK SPACE}", "POPIS"),
    ],
)
def test_header_aliases_case_and_accents(tmp_path: Path, headers: tuple[str, ...]) -> None:
    path = write_xlsx(tmp_path / "valid.xlsx", headers, [("S.11001", "Veřejné", "popis")])
    table = read_table(path, VALID_COLUMNS)
    assert table.rows == ({"Kód": "S.11001", "Název": "Veřejné", "Popis": "popis"},)
    assert set(table.columns) == {"Kód", "Název", "Popis"}


def test_lowercase_value_header_and_extra_columns(tmp_path: Path) -> None:
    path = write_xlsx(
        tmp_path / "cts.xlsx",
        ("Poznámka", "id", "value", "description", "Extra"),
        [("note", 1001, "S.11001", "x", "ignored")],
    )
    table = read_table(path, CTS_COLUMNS)
    assert table.rows == ({"ID": "1001", "VALUE": "S.11001", "DESCRIPTION": "x"},)
    assert table.columns == {"ID": "id", "VALUE": "value", "DESCRIPTION": "description"}


def test_numeric_cells_are_normalized(tmp_path: Path) -> None:
    path = write_xlsx(
        tmp_path / "cts.xlsx",
        ("ID", "VALUE", "DESCRIPTION"),
        [(1001.0, 11001, None), (1002, 1.0, True), ("1003", " 62 ", 3.5)],
    )
    table = read_table(path, CTS_COLUMNS)
    assert table.rows == (
        {"ID": "1001", "VALUE": "11001", "DESCRIPTION": ""},
        {"ID": "1002", "VALUE": "1", "DESCRIPTION": "TRUE"},
        {"ID": "1003", "VALUE": "62", "DESCRIPTION": "3.5"},
    )


def test_blank_rows_are_skipped_but_partial_rows_kept(tmp_path: Path) -> None:
    path = write_xlsx(
        tmp_path / "cts.xlsx",
        ("ID", "VALUE", "DESCRIPTION", "OTHER"),
        [
            (1001, "S.11001", "x", None),
            (None, None, None, "only extra column"),
            (None, "", " ", None),
            (None, None, "description only", None),
            (1002, "S.11002", "y", None),
        ],
    )
    table = read_table(path, CTS_COLUMNS)
    assert table.row_numbers == (2, 5, 6)
    assert table.rows[1] == {"ID": "", "VALUE": "", "DESCRIPTION": "description only"}


def test_named_sheet_is_used(tmp_path: Path) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Info"
    first.append(["This sheet has no table"])
    data = workbook.create_sheet("Data")
    data.append(["ID", "VALUE", "DESCRIPTION"])
    data.append([1, "S.14", "Domácnosti"])
    path = tmp_path / "multi.xlsx"
    workbook.save(path)

    table = read_table(path, CTS_COLUMNS, sheet="Data")
    assert table.sheet == "Data"
    assert table.rows == ({"ID": "1", "VALUE": "S.14", "DESCRIPTION": "Domácnosti"},)
    with pytest.raises(CodebookSchemaError, match="Info"):
        read_table(path, CTS_COLUMNS)
    with pytest.raises(CodebookSchemaError, match="Missing"):
        read_table(path, CTS_COLUMNS, sheet="Missing")


def test_missing_column_raises_schema_error_with_context(tmp_path: Path) -> None:
    path = write_xlsx(
        tmp_path / "CTS_BA0036_NEW.xlsx",
        ("ID", "CODE", "DESCRIPTION"),
        [(1001, "S.11001", "x")],
        title_rows=[("Export",)],
    )
    with pytest.raises(CodebookSchemaError) as info:
        read_table(path, CTS_COLUMNS)
    message = str(info.value)
    assert "CTS_BA0036_NEW.xlsx" in message
    assert "Sheet1" in message
    assert "VALUE" in message and "HODNOTA" in message
    assert "'Export'" in message and "'CODE'" in message


def test_header_beyond_scan_window_is_not_found(tmp_path: Path) -> None:
    path = write_xlsx(
        tmp_path / "cts.xlsx",
        ("ID", "VALUE", "DESCRIPTION"),
        [(1001, "S.11001", "x")],
        title_rows=[("title",)] * 3,
    )
    assert read_table(path, CTS_COLUMNS, header_scan_rows=4).header_row == 4
    with pytest.raises(CodebookSchemaError):
        read_table(path, CTS_COLUMNS, header_scan_rows=3)


def test_alias_matching_two_logical_columns_on_one_cell_is_not_a_header(tmp_path: Path) -> None:
    columns = {"A": ("ID",), "B": ("ID",)}
    path = write_xlsx(tmp_path / "x.xlsx", ("ID", "OTHER"), [(1, 2)])
    with pytest.raises(CodebookSchemaError):
        read_table(path, columns)


def test_missing_file_raises_file_error(tmp_path: Path) -> None:
    with pytest.raises(CodebookFileError, match="not found"):
        read_table(tmp_path / "nope.xlsx", CTS_COLUMNS)


def test_non_xlsx_file_raises_file_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_text("ID;VALUE;DESCRIPTION\n1;S.11001;x\n", encoding="utf-8")
    with pytest.raises(CodebookFileError) as info:
        read_table(path, CTS_COLUMNS)
    assert info.value.__cause__ is not None


def test_empty_columns_mapping_is_rejected(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "x.xlsx", ("ID",), [(1,)])
    with pytest.raises(ValueError):
        read_table(path, {})


def test_header_only_sheet_gives_no_rows(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "empty.xlsx", ("ID", "VALUE", "DESCRIPTION"), [])
    table = read_table(path, CTS_COLUMNS)
    assert table.rows == ()
    assert table.row_numbers == ()


def test_table_rejects_mismatched_row_numbers(tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        Table(
            rows=({"ID": "1"},), row_numbers=(), sheet="S", header_row=1, columns={}, path=tmp_path
        )


@pytest.mark.parametrize("truncate", [True, False])
def test_corrupted_sheet_xml_raises_file_error(tmp_path: Path, truncate: bool) -> None:
    path = write_xlsx(tmp_path / "c.xlsx", ("ID", "VALUE", "DESCRIPTION"), [(1, "S.14", "x")] * 50)
    corrupt_sheet_xml(path, truncate=truncate)
    with pytest.raises(CodebookFileError) as info:
        read_table(path, CTS_COLUMNS)
    assert info.value.__cause__ is not None
    assert "c.xlsx" in str(info.value)
    path.unlink()  # workbook was closed despite the failure


def test_read_table_always_closes_workbook(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import core.codebooks.xlsx as xlsx_module

    closed: list[bool] = []
    real_load = xlsx_module.openpyxl.load_workbook

    def spying_load(*args: object, **kwargs: object) -> object:
        workbook = real_load(*args, **kwargs)
        real_close = workbook.close

        def close() -> None:
            closed.append(True)
            real_close()

        workbook.close = close
        return workbook

    monkeypatch.setattr(xlsx_module.openpyxl, "load_workbook", spying_load)
    good = write_xlsx(tmp_path / "good.xlsx", ("ID", "VALUE", "DESCRIPTION"), [(1, "S.14", "x")])
    read_table(good, CTS_COLUMNS)
    assert closed == [True]
    bad = write_xlsx(tmp_path / "bad.xlsx", ("A", "B", "C"), [(1, 2, 3)])
    with pytest.raises(CodebookSchemaError):
        read_table(bad, CTS_COLUMNS)
    assert closed == [True, True]
    good.unlink()  # PermissionError on Windows if a handle leaked
    bad.unlink()
